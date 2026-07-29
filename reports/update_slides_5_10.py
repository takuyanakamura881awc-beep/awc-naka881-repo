# -*- coding: utf-8 -*-
"""
事前相談会② 状況報告資料のスライド5〜8を再構成し、スライド9・10を追加する。

入力
  - 既存資料（お客様編集済み）: base_in.pptx
  - 構造化データ: reports/data/slide5_8_data_v2.json

方針（指示書 v2 準拠）
  - Slide 5 : 各チームの状況一覧（進捗／10月見込みS-C／難易度／ガイドライン適合性）
  - Slide 6 : エグゼクティブサマリ（S/A/B/C の分布とフォロー優先順位）
  - Slide 7 : 使用アプリ想定と構築ボリューム（相対値。Copilot Studio も計上）
  - Slide 8 : 難易度とガイドライン適合性の分離評価
  - Slide 9 : 事務局判断待ち・横断課題一覧（新規）
  - Slide 10: 8月デモに向けた進め方の型（新規）
  ※ Power Automate は「When an agent calls the flow」のみという制約を全面に反映する
  ※ 既存のチーム個票は 10月見込みの表記のみ S/A/B/C に統一する
"""
import json

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

SRC = ("/tmp/claude-0/-home-user-awc-naka881-repo/"
       "e4943cfb-89be-5c9f-b59c-004830a57633/scratchpad/base_in.pptx")
DATA = "/home/user/awc-naka881-repo/reports/data/slide5_8_data_v2.json"
OUT = ("/home/user/awc-naka881-repo/reports/"
       "CopilotStudio_ハッカソン_事前相談会②_状況報告_20260729.pptx")

# ---- テーマ色（テンプレート準拠）----
DK1 = RGBColor(0x23, 0x23, 0x23)
DK2 = RGBColor(0x1B, 0x5C, 0x80)
LT1 = RGBColor(0xF8, 0xFD, 0xFA)
AC1 = RGBColor(0xEC, 0x85, 0x6F)
AC2 = RGBColor(0x0F, 0xA4, 0xCC)
AC3 = RGBColor(0xA6, 0xE2, 0xF3)
AC4 = RGBColor(0xDB, 0xF1, 0xF7)
AC5 = RGBColor(0xE7, 0xE8, 0xF8)
AC6 = RGBColor(0xFA, 0xE5, 0xE3)
GRAY = RGBColor(0xD0, 0xD4, 0xD7)
MUTE = RGBColor(0x6E, 0x76, 0x7C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# ---- 文字サイズ（テンプレートの体系）----
SZ_XL, SZ_HEAD, SZ_LEAD = 18.14, 16.33, 14.52
SZ_BODY, SZ_FINE = 12.7, 10.89

CX, CW = Inches(0.81), Inches(11.72)
PGX, PGY = Inches(12.67), Inches(7.00)
LY_BODY = 17

prs = Presentation(SRC)
D = json.load(open(DATA, encoding="utf-8"))
TEAMS = {t["team_id"]: t for t in D["teams"]}
ORDER = [f"T{i:02d}" for i in range(1, 12)]


# ==================================================================
# 汎用ヘルパー
# ==================================================================
def txt(s, x, y, w, h, text, size=SZ_BODY, color=DK1, bold=False,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, ls=1.0):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(2); tf.margin_right = Pt(2)
    tf.margin_top = Pt(1); tf.margin_bottom = Pt(1)
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = ls
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return tb


def box(s, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    sp = s.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.fill.background(); sp.shadow.inherit = False
    return sp


def chip(s, x, y, w, h, text, fill, tcolor=DK1, size=SZ_FINE, bold=True):
    sp = box(s, x, y, w, h, fill, MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = sp.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(3); tf.margin_right = Pt(3)
    tf.margin_top = Pt(1); tf.margin_bottom = Pt(1)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = tcolor
    return sp


def cell(c, text, size=SZ_FINE, color=DK1, bold=False, fill=None,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    c.margin_left = Pt(4); c.margin_right = Pt(4)
    c.margin_top = Pt(1); c.margin_bottom = Pt(1)
    c.vertical_anchor = anchor
    c.fill.solid(); c.fill.fore_color.rgb = fill if fill else WHITE
    tf = c.text_frame; tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = 0.96
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color


def table(s, x, y, w, rows, cols, widths, hdr_h, row_h):
    shp = s.shapes.add_table(rows, cols, x, y, w, hdr_h + row_h * (rows - 1))
    t = shp.table
    t.first_row = False
    t.horz_banding = False
    for i, cw in enumerate(widths):
        t.columns[i].width = cw
    t.rows[0].height = hdr_h
    for i in range(1, rows):
        t.rows[i].height = row_h
    return t


def new_slide(title, sub=None):
    s = prs.slides.add_slide(prs.slide_layouts[LY_BODY])
    for ph in s.placeholders:
        if ph.placeholder_format.idx == 0:
            ph.text_frame.paragraphs[0].add_run().text = title
        elif ph.placeholder_format.idx == 2:
            if sub:
                r = ph.text_frame.paragraphs[0].add_run(); r.text = sub
                r.font.size = Pt(SZ_BODY); r.font.color.rgb = MUTE
            else:
                ph._element.getparent().remove(ph._element)
    # ページ番号（最後にまとめて振り直す）
    txt(s, PGX, PGY, Inches(0.54), Inches(0.36), "0", size=SZ_FINE, color=MUTE,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return s


def fc_fill(g):
    return {"S": AC3, "A": AC4, "B": AC5, "C": AC6}.get(g, GRAY)


def fit_fill(v):
    return {"○": AC4, "△": AC5, "×": AC6}.get(v, GRAY)


def diff_fill(v):
    return {"1": AC4, "2": AC4, "3": AC5, "4": AC6}.get(v, GRAY)


# ==================================================================
# 表示用の短縮テキスト（一覧表に収めるための要約）
# ==================================================================
SHORT = {
    "T01": dict(work="提出様式の記載漏れ・形式不備をAIが確認",
                issues="AI判定の揺れ／判定基準の粒度／人の確認範囲",
                action="開発手順の型・8〜9月スケジュール・プロンプト例を共有"),
    "T02": dict(work="受付Cエラーの原因・修正・再申請手順を案内",
                issues="FAQ・ガイドの整備／代表エラーの選定",
                action="設計相談用プロンプトの見本を配布"),
    "T03": dict(work="営業データ統合とリスク分析・訪問前サマリ",
                issues="営業システム等の自動取得可否／外部Web不可／情報区分",
                action="使えるデータをM365内／外／外部Webに分類し可否を明文化"),
    "T04": dict(work="市場調査・企業分析・アイデア出しの自動化",
                issues="外部Web参照禁止に抵触／代替ユースケースが未確定",
                action="再設計対象として最優先フォロー。新案確定後に再評価"),
    "T05": dict(work="依頼事項の登録・通知・回答集約・リマインド",
                issues="Forms回答／スケジュール起動が不可／回答状況の保存先",
                action="通常トリガー禁止・Agent Flow限定をチームへ周知"),
    "T06": dict(work="月次フォーキャスト会議の資料作成と会議招集",
                issues="ファイルサーバ参照／スケジュール実行不可／PPT自動生成",
                action="SharePoint/OneDrive格納前提へ変更し、エージェント起動型へ"),
    "T07": dict(work="請求と関連データの照合・不整合チェック",
                issues="ファイル命名規則の維持／判定基準の固定化",
                action="初回メンター定例でMVPの入力・出力例を確認"),
    "T08": dict(work="会議調整・タスク抽出・レビュー観点の蓄積",
                issues="個人情報・社外秘／予定表・会議室の権限／通常トリガー不可",
                action="個人情報・ゲスト情報の扱いと予定表参照可否を判断事項化"),
    "T09": dict(work="レビュー対象の一覧化・依頼・完了確認",
                issues="共有ファイルサーバ／固定文面の自動送信／受信トリガー不可",
                action="人確認の要否・ファイルサーバ代替・リリース判定要件を整理"),
    "T10": dict(work="文字起こしと会議資料から議事録ドラフト生成",
                issues="文字起こしの取得権限／議事録テンプレート",
                action="サンプル素材と評価観点を準備し、実機確認で精度を担保"),
    "T11": dict(work="社内LAN問い合わせの一次受付（FAQ案内）",
                issues="FAQ整備／本番移行／利用ルール・運用保守",
                action="本番移行の申請フローと運用項目テンプレートを提示"),
}
STAGE_SHORT = {
    "T01": "③ 1案確定", "T02": "③ 1案確定", "T03": "③ 1案確定",
    "T04": "－ 再設計中", "T05": "⑤ 効果試算完了", "T06": "③ 1案確定 *",
    "T07": "⑤ 効果試算完了 *", "T08": "④ 設計具体化", "T09": "④ 設計具体化",
    "T10": "④ 設計具体化", "T11": "④ 設計具体化",
}
APPS = ["Copilot Studio", "Power Automate", "SharePoint", "Teams", "Outlook",
        "Forms", "Power BI", "Excel", "Word/PPT", "その他"]
APP_HDR = ["Copilot\nStudio", "Power\nAutomate", "Share\nPoint", "Teams",
           "Outlook", "Forms", "Power\nBI", "Excel", "Word/\nPPT", "その他"]

PA_RULE = ("Power Automate側でエージェントフローを呼び出すときは、必ず"
           "「When an agent calls the flow（エージェントがフローを呼び出したとき）」"
           "のみ利用すること。ほかのトリガーでは従量課金（Copilot Credits消費）が"
           "発生する可能性があるため、必ず守ること。")


# ==================================================================
# Slide 5：各チームの状況一覧
# ==================================================================
s5 = new_slide("各チームの状況一覧",
               "進捗・10月発表の見込み・難易度・ガイドライン適合性を1枚で把握する")
txt(s5, CX, Inches(1.32), CW, Inches(0.44),
    "10月見込み ▶ S：発表成功の可能性が非常に高い／A：見込み高い（軽微な整理でデモ化可）／"
    "B：重点フォロー（制約確認・スコープ再定義が必要）／C：現行案のままでは難しく再設計が必要\n"
    "難易度 ▶ 1〜4（数字が大きいほど設計・検証の負荷が高い）／「－」は現行案が制約に抵触し評価不可　"
    "｜　ガイドライン適合性 ▶ ○ 制約内で実現可能／△ 事務局判断または設計変更が必要／× 制約に抵触",
    size=9.5, color=DK1, ls=1.22)

t = table(s5, CX, Inches(1.86), CW, 12, 8,
          [Inches(0.72), Inches(2.05), Inches(1.12), Inches(0.72), Inches(0.62),
           Inches(0.72), Inches(3.05), Inches(2.72)], Inches(0.42), Inches(0.42))
for c, h in enumerate(["チーム", "取り組み", "進捗", "10月\n見込み", "難易\n度",
                       "適合性", "主な論点", "事務局アクション"]):
    cell(t.cell(0, c), h, size=SZ_FINE, color=WHITE, bold=True, fill=DK2,
         align=PP_ALIGN.CENTER)
for i, tid in enumerate(ORDER, start=1):
    d = TEAMS[tid]; sh = SHORT[tid]
    rf = WHITE if i % 2 else LT1
    cell(t.cell(i, 0), d["team_label"], size=SZ_FINE, color=DK2, bold=True,
         fill=rf, align=PP_ALIGN.CENTER)
    cell(t.cell(i, 1), sh["work"], size=9.5, color=DK1, bold=True, fill=rf)
    cell(t.cell(i, 2), STAGE_SHORT[tid], size=9.5, color=DK2, bold=True, fill=rf,
         align=PP_ALIGN.CENTER)
    cell(t.cell(i, 3), d["forecast"], size=SZ_BODY, color=DK1, bold=True,
         fill=fc_fill(d["forecast"]), align=PP_ALIGN.CENTER)
    cell(t.cell(i, 4), d["difficulty"], size=SZ_BODY, color=DK1, bold=True,
         fill=diff_fill(d["difficulty"]), align=PP_ALIGN.CENTER)
    cell(t.cell(i, 5), d["guideline_fit"], size=SZ_BODY, color=DK1, bold=True,
         fill=fit_fill(d["guideline_fit"]), align=PP_ALIGN.CENTER)
    cell(t.cell(i, 6), sh["issues"], size=9.5, color=DK1, fill=rf)
    cell(t.cell(i, 7), sh["action"], size=9.5, color=DK1, fill=rf)
ny = Inches(1.86) + Inches(0.42) * 12 + Inches(0.06)
box(s5, CX, ny, CW, Inches(0.42), AC5)
txt(s5, CX + Inches(0.2), ny, CW - Inches(0.4), Inches(0.42),
    "＊ チーム6・7は第2回 事前相談会が未実施のため、第1回と企画書をもとにした暫定評価。"
    "チーム4は現行案が制約に抵触しており、難易度は新案の確定後に再評価する。",
    size=9.5, color=DK1, anchor=MSO_ANCHOR.MIDDLE)


# ==================================================================
# Slide 6：エグゼクティブサマリ
# ==================================================================
s6 = new_slide("エグゼクティブサマリ：全体状況とフォロー優先順位",
               "10月発表に向けた見込みの分布と、事務局が判断すべき事項")
box(s6, CX, Inches(1.34), CW, Inches(0.66), DK2)
txt(s6, CX + Inches(0.3), Inches(1.34), CW - Inches(0.6), Inches(0.66),
    D["slides_recommendation"][1]["message"],
    size=SZ_LEAD, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)

grades = [("S", "1チーム", "チーム7", "MVP範囲・入出力・デモ筋書きが明確で自走可能"),
          ("A", "5チーム", "チーム1・2・5・10・11",
           "軽微な設計整理やプロンプト／ナレッジ整備でデモ化できる"),
          ("B", "4チーム", "チーム3・6・8・9",
           "重点フォロー。制約確認・スコープ再定義・権限／データ配置の整理が必要"),
          ("C", "1チーム", "チーム4", "現行案のままでは難しく、再設計または別案化が必要")]
gy, gh = Inches(2.16), Inches(1.42)
gw, gap = Inches(2.85), Inches(0.09)
for i, (g, cnt, teams, note) in enumerate(grades):
    x = CX + (gw + gap) * i
    box(s6, x, gy, gw, gh, fc_fill(g))
    box(s6, x, gy, gw, Inches(0.42), DK2)
    txt(s6, x, gy, gw, Inches(0.42), f"{g}　{cnt}", size=SZ_HEAD, color=WHITE,
        bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s6, x + Inches(0.12), gy + Inches(0.5), gw - Inches(0.24), Inches(0.3),
        teams, size=SZ_BODY, color=DK1, bold=True, align=PP_ALIGN.CENTER)
    txt(s6, x + Inches(0.14), gy + Inches(0.84), gw - Inches(0.28), Inches(0.5),
        note, size=9.5, color=DK1, align=PP_ALIGN.CENTER, ls=1.14)

txt(s6, CX, Inches(3.76), CW, Inches(0.3), "フォローの優先順位",
    size=SZ_LEAD, color=DK2, bold=True)
prio = [("最優先", "チーム4", AC6,
         "現行案の中核が外部Web参照に依存し制約に抵触。難易度ではなく"
         "「現行案では評価不可・再設計」として扱い、新案の確定を最優先で支援する。"),
        ("重点フォロー", "チーム3・6・8・9", AC5,
         "外部データの取得可否、個人情報・権限、ファイル保管場所、"
         "通常トリガーを使わない起動方式の4点を確認し、MVP範囲を絞り込む。"),
        ("順調・伴走のみ", "チーム1・2・5・7・10・11", AC4,
         "進め方の型・プロンプト見本の共有と、8月の「一本通す」デモ化の伴走で足りる。")]
py = Inches(4.12)
for label, teams, col, note in prio:
    box(s6, CX, py, CW, Inches(0.62), col)
    txt(s6, CX + Inches(0.16), py, Inches(1.5), Inches(0.62), label,
        size=SZ_BODY, color=DK1, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    txt(s6, CX + Inches(1.7), py, Inches(2.2), Inches(0.62), teams,
        size=SZ_BODY, color=DK2, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    txt(s6, CX + Inches(4.0), py, Inches(8.4), Inches(0.62), note,
        size=9.5, color=DK1, anchor=MSO_ANCHOR.MIDDLE, ls=1.14)
    py += Inches(0.7)

box(s6, CX, Inches(6.26), CW, Inches(0.58), AC6)
txt(s6, CX + Inches(0.2), Inches(6.26), CW - Inches(0.4), Inches(0.58),
    "【最重要】Power Automateのエージェントフロー呼び出しは"
    "「When an agent calls the flow」のみ。ほかのトリガーは従量課金"
    "（Copilot Credits消費）が発生する可能性があるため使用しない。",
    size=SZ_BODY, color=DK1, bold=True, anchor=MSO_ANCHOR.MIDDLE)


# ==================================================================
# Slide 7：使用アプリ想定と構築ボリューム
# ==================================================================
s7 = new_slide("使用アプリ想定と構築ボリューム（相対値）",
               "アプリの数ではなく、アプリごとの構築量の相対値と合計で比較する")
txt(s7, CX, Inches(1.3), CW, Inches(0.26),
    "数値は工数ではなく、スライド比較用の相対値（各アプリ0〜40程度）。"
    "Copilot Studioは全チームが使用するため省略せず計上している。",
    size=9.5, color=DK1)

vol_rows = []
for tid in ORDER:
    d = TEAMS[tid]
    av = dict(d["apps_volume"])
    if "OneDrive" in av:                      # OneDrive は SharePoint に合算
        av["SharePoint"] = av.get("SharePoint", 0) + av.pop("OneDrive")
    vol_rows.append((tid, av, d["total_volume"]))
max_total = max(v for _, _, v in vol_rows) or 1

t = table(s7, CX, Inches(1.66), CW, 12, 13,
          [Inches(0.82)] + [Inches(0.79)] * 10 + [Inches(0.78), Inches(2.22)],
          Inches(0.56), Inches(0.385))
cell(t.cell(0, 0), "チーム", size=SZ_FINE, color=WHITE, bold=True, fill=DK2,
     align=PP_ALIGN.CENTER)
for j, h in enumerate(APP_HDR):
    cell(t.cell(0, 1 + j), h, size=9.5, color=WHITE, bold=True, fill=DK2,
         align=PP_ALIGN.CENTER)
cell(t.cell(0, 11), "合計", size=SZ_FINE, color=WHITE, bold=True, fill=DK2,
     align=PP_ALIGN.CENTER)
cell(t.cell(0, 12), "構築ボリューム（相対）", size=SZ_FINE, color=WHITE, bold=True,
     fill=DK2, align=PP_ALIGN.CENTER)
for i, (tid, av, total) in enumerate(vol_rows, start=1):
    rf = WHITE if i % 2 else LT1
    cell(t.cell(i, 0), TEAMS[tid]["team_label"], size=SZ_FINE, color=DK2,
         bold=True, fill=rf, align=PP_ALIGN.CENTER)
    if tid == "T04":
        t.cell(i, 1).merge(t.cell(i, 11))
        cell(t.cell(i, 1), "現行案は制約に抵触しており再設計中のため、構築ボリュームは未評価",
             size=9.5, color=DK1, fill=AC6, align=PP_ALIGN.CENTER)
        cell(t.cell(i, 12), "－", size=SZ_FINE, color=MUTE, fill=rf,
             align=PP_ALIGN.CENTER)
        continue
    for j, app in enumerate(APPS):
        v = av.get(app, 0)
        cell(t.cell(i, 1 + j), (str(v) if v else ""), size=SZ_FINE,
             color=DK1, bold=bool(v),
             fill=(AC3 if v >= 30 else (AC4 if v >= 20 else rf)) if v else rf,
             align=PP_ALIGN.CENTER)
    cell(t.cell(i, 11), str(total), size=SZ_BODY, color=DK2, bold=True,
         fill=rf, align=PP_ALIGN.CENTER)
    cell(t.cell(i, 12), "", size=SZ_FINE, fill=rf)

# 合計の横棒（相対値の視覚化）
bar_x = CX + Inches(0.82) + Inches(0.79) * 10 + Inches(0.78) + Inches(0.12)
bar_w_max = Inches(1.98)
for i, (tid, av, total) in enumerate(vol_rows, start=1):
    if tid == "T04":
        continue
    y = Inches(1.66) + Inches(0.56) + Inches(0.385) * (i - 1) + Inches(0.12)
    w = int(bar_w_max * (total / max_total))
    box(s7, bar_x, y, w, Inches(0.14), AC2)

ny = Inches(1.66) + Inches(0.56) + Inches(0.385) * 11 + Inches(0.06)
box(s7, CX, ny, CW, Inches(0.4), AC5)
txt(s7, CX + Inches(0.2), ny, CW - Inches(0.4), Inches(0.4),
    "＊ 空欄は使用想定なし。数値が大きいアプリほど、そのアプリ上で作り込む量が多い。"
    "チーム6のOneDrive（10）はSharePointに合算している。",
    size=9.5, color=DK1, anchor=MSO_ANCHOR.MIDDLE)


# ==================================================================
# Slide 8：難易度とガイドライン適合性の分離評価
# ==================================================================
s8 = new_slide("難易度とガイドライン適合性の分離評価",
               "「作る量（ボリューム）」「作り方の難しさ（難易度）」「制約適合性」を混同しない")

dl = D["constraints"]["difficulty_definition"]
gl = D["constraints"]["guideline_fit_definition"]
txt(s8, CX, Inches(1.3), Inches(5.7), Inches(0.28), "難易度の定義",
    size=SZ_LEAD, color=DK2, bold=True)
dy = Inches(1.62)
for k in ["1", "2", "3", "4", "-"]:
    box(s8, CX, dy, Inches(0.5), Inches(0.46), diff_fill(k))
    txt(s8, CX, dy, Inches(0.5), Inches(0.46), ("－" if k == "-" else k),
        size=SZ_BODY, color=DK1, bold=True, align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE)
    txt(s8, CX + Inches(0.6), dy, Inches(5.1), Inches(0.46), dl[k],
        size=9.5, color=DK1, anchor=MSO_ANCHOR.MIDDLE, ls=1.12)
    dy += Inches(0.5)

gx = CX + Inches(6.02)
txt(s8, gx, Inches(1.3), Inches(5.7), Inches(0.28), "ガイドライン適合性の定義",
    size=SZ_LEAD, color=DK2, bold=True)
gy2 = Inches(1.62)
for k in ["○", "△", "×"]:
    box(s8, gx, gy2, Inches(0.5), Inches(0.46), fit_fill(k))
    txt(s8, gx, gy2, Inches(0.5), Inches(0.46), k, size=SZ_BODY, color=DK1,
        bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s8, gx + Inches(0.6), gy2, Inches(5.1), Inches(0.46), gl[k],
        size=9.5, color=DK1, anchor=MSO_ANCHOR.MIDDLE, ls=1.12)
    gy2 += Inches(0.5)
box(s8, gx, Inches(3.16), Inches(5.7), Inches(0.86), AC6)
txt(s8, gx + Inches(0.18), Inches(3.16), Inches(5.34), Inches(0.86),
    "難易度「－」は、現行案が制約に抵触し案自体の再設計が必要なため、"
    "現案では評価しないことを示す（チーム4）。新案の確定後に再評価する。",
    size=9.5, color=DK1, anchor=MSO_ANCHOR.MIDDLE, ls=1.14)

# 難易度 × 適合性 マトリクス
txt(s8, CX, Inches(4.2), CW, Inches(0.28), "難易度 × ガイドライン適合性の分布",
    size=SZ_LEAD, color=DK2, bold=True)
mx0, my0 = CX + Inches(0.95), Inches(4.56)
colw, rowh = Inches(2.14), Inches(0.66)
cols = ["1", "2", "3", "4", "－"]
rows_ = ["○", "△", "×"]
for j, c in enumerate(cols):
    box(s8, mx0 + colw * j, my0, colw - Inches(0.05), Inches(0.34), DK2)
    txt(s8, mx0 + colw * j, my0, colw - Inches(0.05), Inches(0.34),
        f"難易度 {c}", size=9.5, color=WHITE, bold=True,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
grid = {("○", "3"): ["1", "2", "7", "10", "11"], ("△", "3"): ["5"],
        ("△", "4"): ["3", "6", "8", "9"], ("×", "－"): ["4"]}
for i, r in enumerate(rows_):
    yy = my0 + Inches(0.38) + rowh * i
    box(s8, CX, yy, Inches(0.9), rowh - Inches(0.05), fit_fill(r))
    txt(s8, CX, yy, Inches(0.9), rowh - Inches(0.05), f"適合性 {r}",
        size=9.5, color=DK1, bold=True, align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE)
    for j, c in enumerate(cols):
        x = mx0 + colw * j
        members = grid.get((r, c), [])
        box(s8, x, yy, colw - Inches(0.05), rowh - Inches(0.05),
            LT1 if not members else AC4)
        if members:
            label = "・".join(f"T{m}" for m in members)
            txt(s8, x, yy, colw - Inches(0.05), rowh - Inches(0.05), label,
                size=SZ_BODY, color=DK2, bold=True, align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE)
txt(s8, CX, Inches(6.6), CW, Inches(0.26),
    "※ 構築ボリューム（作る量）はスライド7を参照。ボリュームが大きくても難易度が高いとは限らない。",
    size=9.5, color=DK1)


# ==================================================================
# Slide 9：事務局判断待ち・横断課題一覧
# ==================================================================
s9 = new_slide("事務局判断待ち・横断課題一覧",
               "複数チームに共通し、事務局の判断・案内が必要な事項")
issues = D["slides_recommendation"][4]["issues"]
short_items = {
    "コスト/ライセンス": "エージェントフローの起動は「When an agent calls the flow」のみ。"
                        "ほかのトリガーは従量課金が発生する可能性があるため使用しない。",
    "外部Web/外部API": "Web検索・RSS・外部API・スクレイピングは不可。"
                       "必要な資料は事前に取得しSharePointに置いて参照する。",
    "個人情報/権限": "社員・協力会社・ゲストの予定表、メール、Teams情報を"
                    "Copilot Studio・Power Automateで利用してよいか。",
    "ファイル保管": "共有ファイルサーバではなくSharePoint／OneDriveへの格納を"
                   "原則化できるか。",
    "本番移行": "開発環境から本番環境への移行、申請、承認、リリース判定の要件。",
    "メール送信": "固定文面のTeams／Outlook通知について、人による確認が必要となる条件。",
}
t = table(s9, CX, Inches(1.42), CW, 7, 3,
          [Inches(2.1), Inches(7.3), Inches(2.32)], Inches(0.4), Inches(0.52))
for c, h in enumerate(["カテゴリ", "事務局に判断・案内いただきたい内容", "該当チーム"]):
    cell(t.cell(0, c), h, size=SZ_FINE, color=WHITE, bold=True, fill=DK2,
         align=PP_ALIGN.CENTER)
for i, it in enumerate(issues, start=1):
    rf = AC6 if it["category"] == "コスト/ライセンス" else (WHITE if i % 2 else LT1)
    cell(t.cell(i, 0), it["category"], size=SZ_FINE, color=DK2, bold=True, fill=rf,
         align=PP_ALIGN.CENTER)
    cell(t.cell(i, 1), short_items[it["category"]], size=9.5, color=DK1, fill=rf)
    cell(t.cell(i, 2), it["teams"].replace("T0", "チーム").replace("T1", "チーム1"),
         size=9.5, color=DK1, fill=rf, align=PP_ALIGN.CENTER)

cy = Inches(1.42) + Inches(0.4) + Inches(0.52) * 6 + Inches(0.12)
box(s9, CX, cy, CW, Inches(1.62), AC6)
box(s9, CX, cy, Inches(0.14), Inches(1.62), AC1)
txt(s9, CX + Inches(0.3), cy + Inches(0.08), CW - Inches(0.5), Inches(0.28),
    "【最重要】Power Automate 起動ルール", size=SZ_LEAD, color=DK1, bold=True)
txt(s9, CX + Inches(0.3), cy + Inches(0.4), Inches(7.0), Inches(0.7),
    PA_RULE, size=SZ_FINE, color=DK1, ls=1.18)
txt(s9, CX + Inches(0.3), cy + Inches(1.14), Inches(7.0), Inches(0.24),
    "使用禁止：スケジュール実行／Dataverse更新／Forms回答／手動ボタン／メール受信／"
    "その他の通常トリガー", size=9.5, color=DK1, bold=True)
txt(s9, CX + Inches(7.6), cy + Inches(0.4), Inches(4.0), Inches(0.24),
    "推奨構成", size=9.5, color=DK1, bold=True)
flow = ["Copilot Studio", "When an agent\ncalls the flow", "Power Automate",
        "SharePoint /\nTeams / Outlook"]
fx = CX + Inches(7.6)
for k, lab in enumerate(flow):
    col = AC1 if k == 1 else DK2
    sp = s9.shapes.add_shape(MSO_SHAPE.CHEVRON, fx, cy + Inches(0.68),
                             Inches(1.18), Inches(0.56))
    sp.fill.solid(); sp.fill.fore_color.rgb = col
    sp.line.fill.background(); sp.shadow.inherit = False
    tf = sp.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(6); tf.margin_right = Pt(3)
    for m, line in enumerate(lab.split("\n")):
        p = tf.paragraphs[0] if m == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = line
        r.font.size = Pt(8); r.font.bold = True
        r.font.color.rgb = WHITE if col == DK2 else DK1
    fx += Inches(1.02)
txt(s9, CX + Inches(7.6), cy + Inches(1.34), Inches(4.1), Inches(0.22),
    "メンターは設計レビュー時に開始トリガーを必ず確認する", size=8.5, color=DK1)


# ==================================================================
# Slide 10：8月デモに向けた進め方の型
# ==================================================================
s10 = new_slide("8月デモに向けた進め方の型",
                "各チームへの共通メッセージ：8月は範囲を絞って「一本通す」")
box(s10, CX, Inches(1.34), CW, Inches(0.6), DK2)
txt(s10, CX + Inches(0.3), Inches(1.34), CW - Inches(0.6), Inches(0.6),
    "8月は作り込みではなく、入力→処理→出力を一本通すことを最優先とする。",
    size=SZ_LEAD, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
steps = [("1", "MVPを1本に絞る",
          "複数の機能を並行して作らず、最も効果が見えるユースケースを1つ選ぶ。"),
         ("2", "入力・処理・出力を1例で定義",
          "サンプルの入力ファイルと、期待する出力の形を先に決める。"),
         ("3", "Copilot Studioから開始する",
          "利用者がエージェントに依頼して動き出す、オンデマンド型の構成にする。"),
         ("4", "フローはエージェント呼び出しのみ",
          "Power Automateは「When an agent calls the flow」でのみ起動する。"),
         ("5", "8月は一本通すデモ、9月は発表・運用整理",
          "9月に発表資料・削減効果・運用項目・手順書の棚卸しを行う。")]
sy2, sh2 = Inches(2.12), Inches(0.78)
for num, ttl, note in steps:
    accent = AC1 if num == "4" else AC2
    bg = AC6 if num == "4" else AC4
    box(s10, CX, sy2, CW, sh2 - Inches(0.1), bg)
    box(s10, CX, sy2, Inches(0.62), sh2 - Inches(0.1), accent)
    txt(s10, CX, sy2, Inches(0.62), sh2 - Inches(0.1), num, size=SZ_XL,
        color=DK1, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s10, CX + Inches(0.82), sy2 + Inches(0.06), Inches(4.3), Inches(0.3),
        ttl, size=SZ_BODY, color=DK2, bold=True)
    txt(s10, CX + Inches(0.82), sy2 + Inches(0.36), Inches(10.6), Inches(0.28),
        note, size=9.5, color=DK1)
    sy2 += sh2
box(s10, CX, Inches(6.14), CW, Inches(0.62), AC5)
txt(s10, CX + Inches(0.2), Inches(6.14), CW - Inches(0.4), Inches(0.62),
    "この型は、8月のメンター定例で各チームの進捗を確認する際の共通の観点としても使用する。"
    "Power Automateを使う構成では、開始トリガーが「When an agent calls the flow」に"
    "なっているかを必ず確認する。",
    size=9.5, color=DK1, anchor=MSO_ANCHOR.MIDDLE, ls=1.16)


# ==================================================================
# 既存スライドの入れ替え・並べ替え
# ==================================================================
RID = qn("r:id")
lst = prs.slides._sldIdLst
ids = list(lst)
for i in sorted([4, 5, 6, 7], reverse=True):          # 旧スライド5〜8を削除
    prs.part.drop_rel(ids[i].get(RID))
    lst.remove(ids[i])
ids = list(lst)
new_ids = ids[-6:]                                     # 追加した6枚を5枚目以降へ移動
for el in new_ids:
    lst.remove(el)
for k, el in enumerate(new_ids):
    lst.insert(4 + k, el)

# ==================================================================
# 各チーム個票：10月見込みの表記を S/A/B/C に統一
# ==================================================================
label_to_id = {TEAMS[t]["team_label"]: t for t in TEAMS}
for s in prs.slides:
    title = ""
    for ph in s.placeholders:
        if ph.placeholder_format.idx == 0:
            title = ph.text_frame.text
            break
    if "：" not in title or not title.startswith("チーム"):
        continue
    tid = label_to_id.get(title.split("：")[0])
    if tid is None:
        continue
    grade = TEAMS[tid]["forecast"]
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        tx = sh.text_frame.text
        if tx.startswith("10月見込み"):
            sh.fill.solid(); sh.fill.fore_color.rgb = fc_fill(grade)
            p0 = sh.text_frame.paragraphs[0]
            for r in list(p0.runs[1:]):
                r._r.getparent().remove(r._r)
            r = p0.runs[0]
            r.text = f"10月見込み {grade}"
            r.font.color.rgb = DK1
        elif tx.startswith("10月発表の見込み（"):
            p0 = sh.text_frame.paragraphs[0]
            base = p0.runs[0]
            for r in list(p0.runs[1:]):
                r._r.getparent().remove(r._r)
            head, _, rest = tx.partition("）")
            base.text = f"10月発表の見込み（{grade}）" + rest

# ==================================================================
# 目次（スライド2）の更新
# ==================================================================
toc_items = [("01", "本日の目的", "第2回の位置づけと、確認・整理する範囲"),
             ("02", "各チームの状況一覧", "進捗・10月見込み・難易度・ガイドライン適合性"),
             ("03", "エグゼクティブサマリ", "S/A/B/Cの分布とフォロー優先順位"),
             ("04", "使用アプリ想定と構築ボリューム", "アプリ別の構築量を相対値で比較"),
             ("05", "難易度とガイドライン適合性", "作る量・難しさ・制約適合性を分けて評価"),
             ("06", "事務局判断待ち・横断課題", "事務局の判断・案内が必要な6項目"),
             ("07", "8月デモに向けた進め方の型", "共通メッセージ：8月は一本通す"),
             ("08", "各チームの状況", "チーム1〜11の個別状況（1チーム1枚）")]
s2 = prs.slides[1]
for sh in list(s2.shapes):
    if sh.is_placeholder and sh.placeholder_format.idx in (0, 2):
        continue
    if sh.has_text_frame and sh.text_frame.text.strip() == "2":
        continue
    sh._element.getparent().remove(sh._element)
ty = Inches(1.5)
for num, ttl, note in toc_items:
    box(s2, CX, ty, Inches(0.56), Inches(0.56), AC4)
    txt(s2, CX, ty, Inches(0.56), Inches(0.56), num, size=SZ_BODY, color=DK2,
        bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s2, CX + Inches(0.76), ty - Inches(0.02), Inches(5.4), Inches(0.32),
        ttl, size=SZ_HEAD, color=DK1, bold=True)
    txt(s2, CX + Inches(0.76), ty + Inches(0.3), Inches(9.8), Inches(0.26),
        note, size=SZ_FINE, color=MUTE)
    ty += Inches(0.66)

# ==================================================================
# ページ番号の振り直し
# ==================================================================
for i, s in enumerate(prs.slides, start=1):
    for sh in s.shapes:
        if not sh.has_text_frame or sh.left is None:
            continue
        if sh.left > Inches(12.4) and sh.top is not None and sh.top > Inches(6.8):
            t0 = sh.text_frame.text.strip()
            if t0.isdigit():
                p0 = sh.text_frame.paragraphs[0]
                for r in list(p0.runs[1:]):
                    r._r.getparent().remove(r._r)
                p0.runs[0].text = str(i)

prs.save(OUT)
print("saved:", OUT)
print("slides:", len(prs.slides._sldIdLst))
