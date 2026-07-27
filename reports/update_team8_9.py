# -*- coding: utf-8 -*-
"""
第2回 事前相談会（チーム8・9 実施分）の内容を、お客様編集済みの資料に反映する。
お客様が加筆された箇所を保持するため、資料を再生成せず既存ファイルを直接更新する。
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

SRC = ("/tmp/claude-0/-home-user-awc-naka881-repo/"
       "e4943cfb-89be-5c9f-b59c-004830a57633/scratchpad/user.pptx")
OUT = ("/home/user/awc-naka881-repo/reports/"
       "CopilotStudio_ハッカソン_事前相談会②_状況報告_20260727.pptx")

DK1 = RGBColor(0x23, 0x23, 0x23)
DK2 = RGBColor(0x1B, 0x5C, 0x80)
AC2 = RGBColor(0x0F, 0xA4, 0xCC)
AC6 = RGBColor(0xFA, 0xE5, 0xE3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTE = RGBColor(0x6E, 0x76, 0x7C)

prs = Presentation(SRC)


def shape_by_id(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    return None


def set_text(sh, text, size=None, color=None, bold=None):
    """書式を保ったままテキストを差し替える（単一段落）"""
    tf = sh.text_frame
    p0 = tf.paragraphs[0]
    for para in list(tf.paragraphs[1:]):
        para._p.getparent().remove(para._p)
    if not p0.runs:
        r = p0.add_run()
    else:
        for r_ in list(p0.runs[1:]):
            r_._r.getparent().remove(r_._r)
        r = p0.runs[0]
    r.text = text
    if size is not None:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    if bold is not None:
        r.font.bold = bold


def set_fill(sh, color):
    sh.fill.solid()
    sh.fill.fore_color.rgb = color


def add_bullet(slide, x, y, w, text, marker_x, marker_y, marker_color,
               size=10.89):
    """箇条書き（マーカー＋本文）を追加"""
    mk = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, marker_x, marker_y,
                                Inches(0.1), Inches(0.1))
    set_fill(mk, marker_color)
    mk.line.fill.background()
    mk.shadow.inherit = False
    tb = slide.shapes.add_textbox(x, y, w, Inches(0.28))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(2); tf.margin_right = Pt(2)
    tf.margin_top = Pt(1); tf.margin_bottom = Pt(1)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = DK1
    return tb


# ==================================================================
# S5 各チームの状況一覧（index 4）
# ==================================================================
s = prs.slides[4]
set_text(shape_by_id(s, 3),
         "全11チームが1案に確定／第2回 実施済み＝チーム6・7以外の9チーム",
         size=12.7, color=MUTE)

for sh in s.shapes:
    if sh.has_text_frame and "進捗の段階" in sh.text_frame.text:
        sh.height = Inches(0.6)

tbl = None
for sh in s.shapes:
    if sh.has_table:
        tbl = sh.table
        break
# 第2回の主な論点（列3）とチーム8・9の文字色を更新
rows = {8: "Outlook・Teams参照時の個人情報／社外秘の取り扱い（事務局判断を要望）",
        9: "定型メールの自動送信時に人の確認が必要か／チーム内の分業の進め方"}
for team_no, topic in rows.items():
    c = tbl.cell(team_no, 3)
    set_text(c, topic, size=10.89, color=DK1)

# チーム9の進捗を ③1案確定 → ④設計具体化 に更新（4本目の矢印を点灯）
row_y = 1.94 + 0.38 + 0.418 * 8 + 0.418 / 2      # 9行目の中心 y(inch)
chevs, labels = [], []
for sh in s.shapes:
    if sh.top is None:
        continue
    yc = (sh.top + sh.height / 2) / 914400
    if abs(yc - row_y) > 0.06:
        continue
    if sh.shape_type is not None and "AUTO_SHAPE" in str(sh.shape_type) \
            and sh.width < Inches(0.35):
        chevs.append(sh)
    if sh.has_text_frame and "1案確定" in sh.text_frame.text:
        labels.append(sh)
chevs.sort(key=lambda x: x.left)
if len(chevs) >= 4:
    set_fill(chevs[3], AC2)
for lb in labels:
    set_text(lb, "④ 設計具体化", size=10.89, color=DK2, bold=True)

# ==================================================================
# S8 第2回で挙がった課題（index 7）
# ==================================================================
s = prs.slides[7]
set_text(shape_by_id(s, 3),
         "チーム1・2・3・4・5・8・9・10・11で実際に挙がった論点を集約",
         size=12.7, color=MUTE)

AC1 = RGBColor(0xEC, 0x85, 0x6F)
# --- カード01：既存2件を詰めて、事務局確認事項2件を追加 ---
b1, m1 = shape_by_id(s, 11), shape_by_id(s, 10)
b2, m2 = shape_by_id(s, 13), shape_by_id(s, 12)
act1 = shape_by_id(s, 14)
for sh in (b1, b2):
    for para in sh.text_frame.paragraphs:
        for r in para.runs:
            r.font.size = Pt(10.89)
b1.top, b1.height = Inches(1.54), Inches(0.44)
m1.top = Inches(1.62)
b2.top, b2.height = Inches(2.00), Inches(0.28)
m2.top = Inches(2.08)
add_bullet(s, Inches(4.75), Inches(2.28), Inches(7.60),
           "社員・協力会社・ゲストの氏名／メールをCopilot Studio等で"
           "利用してよいか　… チーム8・9",
           Inches(4.53), Inches(2.36), AC1)
add_bullet(s, Inches(4.75), Inches(2.56), Inches(7.60),
           "Power Automateで固定文面を自動送信する場合も人の確認が必要か"
           "（定型通知の条件）　… チーム9",
           Inches(4.53), Inches(2.64), AC1)
act1.top = Inches(2.84)

# --- カード03：既存3件を詰めて、リリース要件を追加 ---
c3 = [(shape_by_id(s, 33), shape_by_id(s, 32)),
      (shape_by_id(s, 35), shape_by_id(s, 34)),
      (shape_by_id(s, 37), shape_by_id(s, 36))]
for i, (b, m) in enumerate(c3):
    for para in b.text_frame.paragraphs:
        for r in para.runs:
            r.font.size = Pt(10.89)
    b.top, b.height = Inches(5.28 + 0.30 * i), Inches(0.28)
    m.top = Inches(5.36 + 0.30 * i)
add_bullet(s, Inches(4.75), Inches(6.18), Inches(7.60),
           "リリース判定に必要な要件（必須テスト・承認プロセス・セキュリティ確認）の共有"
           "　… チーム9",
           Inches(4.53), Inches(6.26), AC2)
act3 = shape_by_id(s, 38)
act3.top = Inches(6.52)
set_text(act3, "▶ 事務局へ：展開・運用の枠組み、定例の進め方、リリース判定の要件を案内いただきたい。",
         size=12.7, color=DK2, bold=True)

# 3件の事務局確認事項も含む旨をカード02の助言行に触れない（カード01/03に反映済み）

# ==================================================================
# S17 チーム8（index 16）
# ==================================================================
s = prs.slides[16]
chip = shape_by_id(s, 15)
set_fill(chip, DK2)
set_text(chip, "第2回 実施済", size=12.7, color=WHITE, bold=True)
set_text(shape_by_id(s, 17),
         "10月発表の見込み（○）の根拠：会議調整とタスク整理の構成は明確。"
         "個人情報の取り扱いに関する事務局判断が前提となるが、範囲を絞れば発表できる。",
         size=10.89, color=DK1)
set_text(shape_by_id(s, 41),
         "予定表・会議室・メールへのアクセス権限に加え、参照範囲に含まれる"
         "個人情報・社外秘の取り扱い可否（事務局判断）が前提となる。",
         size=10.89, color=DK1)
set_text(shape_by_id(s, 45),
         "第2回：会議室予約・タスク管理でOutlook・Teamsを参照するため、"
         "個人情報や社外秘が含まれる。取り扱いを事務局で決めてほしい。"
         "あわせて外部Web参照が禁止となる理由も確認。",
         size=10.89, color=DK1)
set_text(shape_by_id(s, 49),
         "参照範囲には社員・協力会社・ゲスト（当社メンバーを含む）の個人情報が入りうるため、"
         "事務局判断として整理し回答する。外部Web参照はスクレイピングに該当し、"
         "相手先サーバへの負荷や攻撃と見なされるリスクがあるため原則不可。",
         size=10.89, color=DK1)
set_text(shape_by_id(s, 52),
         "個人情報の取り扱い可否について事務局の回答を待ちつつ、"
         "会議調整とタスク抽出のどちらを中心にするか決め、デモの筋書きを1本作る。",
         size=10.89, color=DK1)

# ==================================================================
# S18 チーム9（index 17）
# ==================================================================
s = prs.slides[17]
ids = {sh.shape_id: sh for sh in s.shapes}
# チップ（第2回 未実施 → 実施済）※チップは幅2インチ未満の図形に限定して特定する
for sh in s.shapes:
    if (sh.shape_type is not None and "AUTO_SHAPE" in str(sh.shape_type)
            and sh.width < Inches(2.0) and sh.has_text_frame
            and "第2回" in sh.text_frame.text):
        set_fill(sh, DK2)
        set_text(sh, "第2回 実施済", size=12.7, color=WHITE, bold=True)
# 進捗 ③1案確定 → ④設計まで具体化（4本目の矢印を点灯）
chevs = [sh for sh in s.shapes
         if sh.shape_type is not None and "AUTO_SHAPE" in str(sh.shape_type)
         and sh.width < Inches(0.35) and abs(sh.top / 914400 - 1.59) < 0.08]
chevs.sort(key=lambda x: x.left)
if len(chevs) >= 4:
    set_fill(chevs[3], AC2)
for sh in s.shapes:
    if sh.has_text_frame and "1案確定" in sh.text_frame.text:
        set_text(sh, "④ 設計まで具体化", size=12.7, color=DK2, bold=True)
# 本文の更新（テキスト内容で対象を特定）
def find_by(slide, keyword):
    for sh in slide.shapes:
        if sh.has_text_frame and keyword in sh.text_frame.text:
            return sh
    return None


set_text(find_by(s, "10月発表の見込み"),
         "10月発表の見込み（○）の根拠：作業を3フェーズに分解し分業方針も定まった。"
         "進め方は自走できており、事務局確認（メール送信・ファイルサーバー）が前提。",
         size=10.89, color=DK1)
set_text(find_by(s, "開発中案件・社外秘情報の扱い"),
         "定型メールの自動送信で人の確認が必要かの判断、共有ファイルサーバーの参照可否が前提。"
         "社外秘情報の扱いも整理が必要。",
         size=10.89, color=DK1)
set_text(find_by(s, "第2回は未実施"),
         "第2回：Power Automateで固定文面を自動送信する場合も人の確認が必要か。"
         "ファイルサーバーの扱い。チーム内で分業が可能か。機能の使い方も習得したい。",
         size=10.89, color=DK1)
set_text(find_by(s, "管理ツールとしての性格"),
         "定型文書であれば確認不要と考えられるが、送信条件は事務局に確認する。"
         "作業は①Automateでのファイル収集②Copilot Studioでのチェック③Automateでの"
         "タスク管理の3フェーズに分かれるため、各工程の入出力を定義して分担し、"
         "1名が全体管理を行う。分担案はCopilotチャットに各自の理解度を入力して相談するとよい。",
         size=10.89, color=DK1)
set_text(find_by(s, "まずメール検知"),
         "3フェーズ（収集→チェック→タスク管理）で担当を分け、各工程の入力・出力を定義する。"
         "ファイルサーバーはSharePointでどこまで対応できるかを確認し、"
         "Copilot Studio研修で機能を習得する。",
         size=10.89, color=DK1)

prs.save(OUT)
print("saved:", OUT)
print("slides:", len(prs.slides._sldIdLst))
