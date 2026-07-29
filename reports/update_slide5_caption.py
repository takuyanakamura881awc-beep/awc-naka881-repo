# -*- coding: utf-8 -*-
"""スライド5「各チームの状況一覧」の 進捗／10月見込み／難易度 の各セルに
評価値の下へ小さな説明文（キャプション）を追加する。"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

F = "/home/user/awc-naka881-repo/reports/CopilotStudio_ハッカソン_事前相談会②_状況報告_20260729.pptx"

CAP_SIZE = Pt(7)
CAP_COLOR = RGBColor(0x59, 0x59, 0x59)

# 進捗レベル → 説明
STEP_CAP = {
    "①": "案を整理中",
    "②": "1案化済み",
    "③": "MVP・入出力OK",
    "④": "概念設計まで",
    "⑤": "効果・筋書きまで",
    "－": "再設計中",
    "-": "再設計中",
}
# 10月見込み → 説明
FC_CAP = {
    "S": "可能性大",
    "A": "見込み高い",
    "B": "重点フォロー",
    "C": "現行案では難",
}
# 難易度 → 説明
DF_CAP = {
    "1": "低い",
    "2": "やや低い",
    "3": "標準",
    "4": "高い",
    "－": "評価不可",
    "-": "評価不可",
}

COL_STEP, COL_FC, COL_DF = 2, 3, 4
SHIFT = Inches(0.06)   # ステップインジケーターを上へ寄せる量


def style_cap(run):
    run.font.size = CAP_SIZE
    run.font.bold = False
    run.font.color.rgb = CAP_COLOR


def add_caption_para(cell, text):
    """既存の評価値パラグラフの下に、小さな説明パラグラフを追加"""
    tf = cell.text_frame
    base = tf.paragraphs[0]
    p = copy.deepcopy(base._p)
    base._p.getparent().append(p)
    from pptx.text.text import _Paragraph
    para = _Paragraph(p, tf)
    for r in list(para.runs):
        r._r.getparent().remove(r._r)
    run = para.add_run()
    run.text = text
    style_cap(run)
    para.alignment = PP_ALIGN.CENTER
    para.space_before = Pt(1)
    para.space_after = Pt(0)
    return para


def main():
    prs = Presentation(F)
    s = prs.slides[4]

    tbl = [sh for sh in s.shapes if sh.has_table][0].table

    # --- 10月見込み・難易度：値の下にキャプションを追加 ---
    for i in range(1, 12):
        for col, cmap in ((COL_FC, FC_CAP), (COL_DF, DF_CAP)):
            cell = tbl.cell(i, col)
            val = cell.text_frame.paragraphs[0].runs[0].text.strip()
            cap = cmap.get(val)
            if cap:
                add_caption_para(cell, cap)

    # --- 進捗：ステップインジケーターを上に寄せ、セル下端にキャプション ---
    # 進捗レベルを表す小さなテキストボックス（○付き数字）を行順に取得
    lvl_boxes = sorted(
        [sh for sh in s.shapes
         if sh.shape_type == 17 and sh.has_text_frame
         and sh.left > Inches(5.2) and sh.left < Inches(5.5)
         and Inches(2.0) < sh.top < Inches(7.0)],
        key=lambda sh: sh.top)
    chevrons = [sh for sh in s.shapes if sh.name.startswith("Chevron")]

    for sh in lvl_boxes + chevrons:
        sh.top = sh.top - SHIFT

    assert len(lvl_boxes) == 11, len(lvl_boxes)

    for i, box in enumerate(lvl_boxes, start=1):
        mark = box.text_frame.text.replace("*", "").replace("＊", "").strip()
        cap = STEP_CAP.get(mark)
        if not cap:
            print("  ! 進捗キャプション未定義:", repr(mark))
            continue
        cell = tbl.cell(i, COL_STEP)
        tf = cell.text_frame
        cell.vertical_anchor = MSO_ANCHOR.BOTTOM
        para = tf.paragraphs[0]
        for r in list(para.runs):
            r._r.getparent().remove(r._r)
        run = para.add_run()
        run.text = cap
        style_cap(run)
        para.alignment = PP_ALIGN.CENTER

    prs.save(F)
    print("saved", F)


if __name__ == "__main__":
    main()
