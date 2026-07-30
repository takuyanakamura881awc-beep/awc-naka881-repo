# -*- coding: utf-8 -*-
"""メンター試作用 チーム別Excel詳細設計書（1チーム1ファイル）を生成する。

入力（正本）
- priority_medium_excel_design_input_v1.zip
    common/excel_template_spec.md, sheet_definition.md, design_style_guide.md
    data/Txx_design_data.json, images/Txx_architecture.png
- priority_medium_excel_design_input_v3_high_quality_visuals.zip
    common/claude_generation_prompt_visual_update_v3.md
    images/team_image_counts.json, images/high_quality_visual_manifest.json
    ＋ 各チームの追加画像50点（v2は使用しない）
- 先に納品した priority_medium_mentor_design（V1）の設計データ
    JSONに無い共通シート（03/06/07/12）と補足に転記して使用する

方針
- 共通14シートは削除しない。設計データが無い場合も枠と理由を残す。
- アプリ固有シートは JSON の app_specific_sheets に挙がっているものだけ作成する。
- 画像は原寸のまま貼り付ける（縮小・再圧縮・簡素化はしない）。
- 画像だけで終わらせず、構築に必要な表データを必ず併記する。
- T04は対象外。00_表紙・前提に除外理由を記載する。
"""
import json
import os

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
SCRATCH = ("/tmp/claude-0/-home-user-awc-naka881-repo/"
           "e4943cfb-89be-5c9f-b59c-004830a57633/scratchpad")
IN = os.path.join(SCRATCH, "xl_in")
MENTOR_V1 = os.path.join(HERE, "..", "mentor_pack", "v1_data.json")
OUTDIR = os.path.join(HERE, "out")

TEAMS = ["T01", "T02", "T03", "T05", "T06",
         "T07", "T08", "T09", "T10", "T11"]
VERSION = "v1.0"
MADE_ON = "2026-07-29"
VENDOR = "株式会社Low Code"

# ---- 配色（Low Codeテーマ ＋ design_style_guide.md の役割） ----
NAVY = "1B5C80"     # 濃紺ヘッダー
TEAL = "0FA4CC"     # セクション見出し（文字は本文色。白抜きは2.9:1で不足）
LTBLUE = "DBF1F7"   # 薄青（副見出し・帯）
LAVEND = "E7E8F8"   # 薄紫（表の縞）
ORANGE = "EC856F"   # 注意
LTRED = "FAE5E3"    # リスク
INK = "232323"      # 静的説明（黒）
GRAY = "595959"     # 静的説明（グレー）
BLUEINK = "1B5C80"  # 入力値（青文字）
GREENINK = "2E7D32"  # 完了/OK
WHITE = "FFFFFF"

THIN = Side(style="thin", color="BFC7CC")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# 追加画像の配置先（v3 claude_generation_prompt_visual_update_v3.md の推奨配置）
IMG_SHEET = {
    "02_mvp_flow": "03_全体業務フロー",
    "02_data_flow": "03_全体業務フロー",
    "03_role_split": "06_人AI自動化分担",
    "03_copilot_studio_design": "CS_CopilotStudio設計",
    "03_knowledge_design": "CS_CopilotStudio設計",
    "03_knowledge_structure": "CS_CopilotStudio設計",
    "03_sharepoint_excel": "SPO_リスト設計",
    "03_excel_sharepoint": "SPO_リスト設計",
    "03_sharepoint_word": "SPO_リスト設計",
    "04_sharepoint_excel_test": "SPO_リスト設計",
    "04_sharepoint_lists": "SPO_リスト設計",
    "04_sharepoint_review": "SPO_リスト設計",
    "04_sharepoint_tasklist": "SPO_リスト設計",
    "04_storage_design": "OneDrive_保管設計",
    "04_copilot_answer": "CS_CopilotStudio設計",
    "04_copilot_difference": "CS_CopilotStudio設計",
    "04_copilot_minutes": "CS_CopilotStudio設計",
    "04_copilot_visit_summary": "CS_CopilotStudio設計",
    "05_copilot_conversation": "CS_CopilotStudio設計",
    "05_copilot_document_design": "CS_CopilotStudio設計",
    "05_copilot_review": "CS_CopilotStudio設計",
    "04_test_escalation": "09_テストケース",
    "05_escalation_log": "SPO_リスト設計",
    "05_risk_constraints": "10_リスク・躓きポイント",
    "05_test_plan": "09_テストケース",
    "06_test_plan": "09_テストケース",
    "07_test_plan": "09_テストケース",
    "08_test_plan": "09_テストケース",
    "05_word_output": "WordPPT_出力設計",
    "07_output_design": "WordPPT_出力設計",
    "06_power_automate_flow": "PA_PowerAutomate設計",
    "06_power_automate_teams": "PA_PowerAutomate設計",
    "06_outlook_boundary": "Outlook_メール予定表設計",
    "07_outlook_boundary": "Outlook_メール予定表設計",
}

IMG_CAPTION = {
    "02_mvp_flow": "MVP全体フロー（入力→AI処理→保管→出力→人の確認）",
    "02_data_flow": "データフロー（人手取得→SharePoint格納→整形→サマリ生成）",
    "03_role_split": "人 / AI / 自動化の役割分担",
    "03_copilot_studio_design": "Copilot Studio エージェント設計",
    "03_knowledge_design": "ナレッジ（FAQ・手順書）設計",
    "03_knowledge_structure": "ナレッジ構成（FAQ・手順・申請種別）",
    "03_sharepoint_excel": "SharePoint / Excel のデータ配置",
    "03_excel_sharepoint": "Excel整形と SharePoint 格納",
    "03_sharepoint_word": "SharePoint / Word のデータ配置",
    "04_sharepoint_excel_test": "SharePoint・Excel・テストの関係",
    "04_sharepoint_lists": "SharePointリスト設計",
    "04_sharepoint_review": "SharePointリスト設計（レビュー管理表）",
    "04_sharepoint_tasklist": "SharePointリスト設計（タスク一覧）",
    "04_storage_design": "ファイル保管設計（OneDrive / SharePoint）",
    "04_copilot_answer": "Copilot Studio 回答生成と切り分け",
    "04_copilot_difference": "Copilot Studio 差異判定",
    "04_copilot_minutes": "Copilot Studio 議事録ドラフト生成",
    "04_copilot_visit_summary": "Copilot Studio 訪問前サマリ生成",
    "05_copilot_conversation": "Copilot Studio 会話設計（アダプティブカード含む）",
    "05_copilot_document_design": "Copilot Studio 資料生成設計",
    "05_copilot_review": "Copilot Studio レビュー観点整理",
    "04_test_escalation": "テストとエスカレーション",
    "05_escalation_log": "エスカレーション記録の管理",
    "05_risk_constraints": "リスクと制約（外部Web禁止・マスキング）",
    "05_test_plan": "テスト計画",
    "06_test_plan": "テスト計画",
    "07_test_plan": "テスト計画",
    "08_test_plan": "テスト計画",
    "05_word_output": "Word 出力設計",
    "07_output_design": "出力設計（PowerPoint / 通知文案）",
    "06_power_automate_flow": "Power Automate フロー設計",
    "06_power_automate_teams": "Power Automate と Teams 通知",
    "06_outlook_boundary": "Outlook 利用範囲の線引き",
    "07_outlook_boundary": "Outlook 利用範囲の線引き",
}

# 設計データが JSON に無いアプリ固有シートの列枠
FRAME_COLS = {
    "Teams_通知設計": (
        ["No", "通知名", "通知契機", "通知先", "通知手段", "メッセージ内容",
         "添付・リンク", "確認者", "注意事項", "メンター確認ポイント"],
        [6, 18, 26, 20, 18, 34, 22, 14, 30, 26]),
    "Outlook_メール予定表設計": (
        ["No", "機能", "対象", "契機", "宛先", "件名テンプレート",
         "本文テンプレート", "送信前確認", "注意事項", "メンター確認ポイント"],
        [6, 18, 14, 24, 18, 28, 34, 20, 30, 26]),
    "Excel_データ処理設計": (
        ["No", "ファイル名", "シート名", "処理内容", "入力範囲", "出力範囲",
         "更新方法", "注意事項", "メンター確認ポイント"],
        [6, 22, 16, 34, 20, 20, 22, 30, 26]),
    "WordPPT_出力設計": (
        ["No", "出力物", "テンプレート", "差し込み項目", "生成方法", "保存先",
         "命名規則", "人の確認", "注意事項"],
        [6, 20, 22, 30, 26, 24, 22, 20, 28]),
    "OneDrive_保管設計": (
        ["No", "フォルダ", "用途", "格納物", "命名規則", "共有範囲",
         "保持方針", "注意事項"],
        [6, 20, 24, 26, 22, 20, 22, 28]),
}

with open(MENTOR_V1, encoding="utf-8") as f:
    MV = json.load(f)
V1T = {t["id"]: t for t in MV["teams"]}
V1C = MV["common"]
CONSTRAINTS = V1C["constraints"]["bullets"]
PRINCIPLES = [ln.split(". ", 1)[1] for ln in V1C["principles"]["raw"].splitlines()
              if ln[:1].isdigit()]
CHECKLIST = V1C["checklist"]["bullets"]
SAMPLE_POLICY = V1C["sample_policy"]["bullets"]
ATTENTION = [(k, " ".join(v))
             for k, v in V1C["attention"]["sections"].items()]
QUALITY_ADD = [a.lstrip("- ").strip()
               for a in V1C["quality"]["追加確認が望ましい点"]]
T04_NOTE = [ln.strip() for ln in V1C["t04"].splitlines()
            if ln.strip() and not ln.startswith("#")]

COMMON_SHEETS = [
    "00_表紙・前提", "01_企画概要", "02_MVPスコープ", "03_全体業務フロー",
    "04_アーキテクチャ", "05_使用アプリ一覧", "06_人AI自動化分担",
    "07_構築ステップ総覧", "08_ダミーデータ一覧", "09_テストケース",
    "10_リスク・躓きポイント", "11_メンター確認チェックリスト",
    "12_未決事項・事務局確認", "13_変更履歴",
]

SHEET_PURPOSE = {
    "00_表紙・前提": "本設計書の対象・前提・共通制約と、収録シートの一覧",
    "01_企画概要": "何のための仕組みか（目的・入力・出力・使用アプリ）",
    "02_MVPスコープ": "8月MVPでどこまで作るか／作らないか",
    "03_全体業務フロー": "入力から人の確認までの工程と主体",
    "04_アーキテクチャ": "システム構成要素と役割",
    "05_使用アプリ一覧": "使用アプリと、使用しないアプリの理由",
    "06_人AI自動化分担": "人・Copilot Studio・Power Automate・保管先の役割",
    "07_構築ステップ総覧": "メンターが手元で試作する手順（全体）",
    "08_ダミーデータ一覧": "試作に使うダミーデータの仕様とサンプル行",
    "09_テストケース": "MVPが一本通ることを確認するテスト",
    "10_リスク・躓きポイント": "想定される躓きと対応案",
    "11_メンター確認チェックリスト": "試作前後にメンターが確認する項目",
    "12_未決事項・事務局確認": "事務局・運営に確認が必要な論点",
    "13_変更履歴": "本設計書の版管理",
    "CS_CopilotStudio設計": "エージェントの目的・指示・ナレッジ・回答ルール",
    "SPO_リスト設計": "SharePointリストの列定義（内部名・種類・必須・入力例）",
    "SPO_ライブラリ設計": "ライブラリのフォルダ・命名規則・権限",
    "PA_PowerAutomate設計": "フローの処理設計（実装前設計レベル・B粒度）",
    "Teams_通知設計": "Teams通知の契機・宛先・文面",
    "Outlook_メール予定表設計": "メール送信・予定表参照の範囲と文面",
    "Excel_データ処理設計": "Excelでの整形・集計処理",
    "WordPPT_出力設計": "Word / PowerPoint 出力物の設計",
    "OneDrive_保管設計": "OneDriveのフォルダ構成と保管方針",
}


# ==================================================================
# Excel ヘルパー
# ==================================================================
class Sheet:
    """1シートを上から順に組み立てるためのカーソル付きラッパー"""

    def __init__(self, wb, name, ctx, widths):
        self.ws = wb.create_sheet(name)
        self.name = name
        self.ctx = ctx
        self.row = 1
        self.freeze_at = None
        self.ncol = len(widths)
        for i, w in enumerate(widths, 1):
            self.ws.column_dimensions[get_column_letter(i)].width = w
        self.ws.sheet_view.showGridLines = False
        self._title()

    # ---- 見出し ----
    def _merge(self, row, height=None):
        self.ws.merge_cells(start_row=row, start_column=1,
                            end_row=row, end_column=self.ncol)
        if height:
            self.ws.row_dimensions[row].height = height

    def _title(self):
        c = self.ctx
        cell = self.ws.cell(1, 1)
        cell.value = "%s　｜　%s %s（%s）" % (
            self.name, c["team_id"], c["team_name"], c["agent_name"])
        cell.font = Font(bold=True, size=14, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", indent=1)
        self._merge(1, 30)
        sub = self.ws.cell(2, 1)
        sub.value = "%s　／　%s　／　%s　／　作成日 %s　／　%s" % (
            SHEET_PURPOSE.get(self.name, ""), c["design_status"],
            "版 " + VERSION, MADE_ON, VENDOR)
        sub.font = Font(size=9, color=GRAY)
        sub.alignment = Alignment(vertical="center", indent=1)
        self._merge(2, 18)
        self.row = 4

    def section(self, text, kind="teal"):
        fill = {"teal": TEAL, "blue": LTBLUE, "warn": ORANGE,
                "risk": LTRED}[kind]
        cell = self.ws.cell(self.row, 1)
        cell.value = "■ " + text
        cell.font = Font(bold=True, size=11, color=INK)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(vertical="center", indent=1)
        self._merge(self.row, 22)
        self.row += 1

    def note(self, text, kind="gray"):
        color = {"gray": GRAY, "warn": INK, "ok": GREENINK,
                 "ink": INK, "risk": INK}[kind]
        cell = self.ws.cell(self.row, 1)
        cell.value = text
        cell.font = Font(size=9, color=color,
                         bold=(kind in ("warn", "ok")))
        cell.alignment = Alignment(vertical="center", wrap_text=False, indent=1)
        if kind == "warn":
            cell.fill = PatternFill("solid", fgColor=ORANGE)
            self._merge(self.row, 20)
        elif kind == "risk":
            cell.fill = PatternFill("solid", fgColor=LTRED)
            self._merge(self.row, 20)
        else:
            self._merge(self.row, 16)
        self.row += 1

    def blank(self, n=1):
        self.row += n

    # ---- 表 ----
    def table(self, headers, rows, ink_cols=(), height=None, filt=True):
        r0 = self.row
        for j, h in enumerate(headers, 1):
            c = self.ws.cell(r0, j)
            c.value = h
            c.font = Font(bold=True, size=10, color=WHITE)
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.alignment = Alignment(horizontal="center", vertical="center",
                                    wrap_text=True)
            c.border = BOX
        self.ws.row_dimensions[r0].height = 26
        if self.freeze_at is None:
            self.freeze_at = r0 + 1
        for i, row in enumerate(rows):
            r = r0 + 1 + i
            for j, v in enumerate(row, 1):
                c = self.ws.cell(r, j)
                c.value = "" if v is None else v
                is_ink = headers[j - 1] in ink_cols
                c.font = Font(size=10,
                              color=BLUEINK if is_ink else INK,
                              bold=False)
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.fill = PatternFill("solid",
                                     fgColor=WHITE if i % 2 == 0 else "F5FAFC")
                c.border = BOX
            if height:
                self.ws.row_dimensions[r].height = height
        if filt and rows:
            self.ws.auto_filter.ref = "A%d:%s%d" % (
                r0, get_column_letter(len(headers)), r0 + len(rows))
        self.row = r0 + len(rows) + 2
        return r0

    def kv(self, pairs, w_label=None, ink_keys=()):
        """項目／内容の2列表"""
        self.table(["項目", "内容"], [[k, v] for k, v in pairs],
                   ink_cols=("内容",) if ink_keys == "all" else (),
                   filt=False)

    # ---- 画像 ----
    def image(self, path, caption, source_note):
        self.section("参考図：" + caption, "blue")
        self.note("画像ファイル： %s　（v3 高品質ビジュアル・原寸のまま貼付／"
                  "再生成・簡略化・再圧縮なし）%s" % (os.path.basename(path),
                                            source_note))
        img = XLImage(path)
        anchor_row = self.row
        img.anchor = "A%d" % anchor_row
        self.ws.add_image(img)
        # 既定行高(15pt=20px)で画像の高さぶんの行を確保する
        self.row = anchor_row + (img.height + 19) // 20 + 2

    def finish(self):
        if self.freeze_at:
            self.ws.freeze_panes = "A%d" % self.freeze_at
        else:
            self.ws.freeze_panes = "A4"


# ==================================================================
# データ準備
# ==================================================================
def load_images():
    idx = {}
    manifest = json.load(open(
        os.path.join(IN, "images", "high_quality_visual_manifest.json"),
        encoding="utf-8"))
    for m in manifest:
        p = os.path.join(IN, m["image"])
        if not os.path.exists(p):
            raise SystemExit("画像が見つかりません: " + m["image"])
        idx.setdefault(m["team_id"], []).append((m["visual_type"], p))
    counts = json.load(open(
        os.path.join(IN, "images", "team_image_counts.json"), encoding="utf-8"))
    for t, n in counts.items():
        if len(idx.get(t, [])) != n:
            raise SystemExit("画像枚数が team_image_counts.json と不一致: " + t)
    return idx


def flow_rows(tid, d):
    """03_全体業務フロー の工程表。V1のMermaid構成と分担から作る"""
    v = V1T[tid]
    out = v["mermaid_out"] or d["overview"]["mvp"]
    return [
        ["1", "入力準備", "人", d["overview"].get("_inputs", ""),
         "入力ファイル・依頼内容をSharePoint／OneDriveに配置する",
         "処理対象データ", "実データ・個人情報を含めない"],
        ["2", "意図理解・判断", "Copilot Studio", "利用者の入力",
         "目的・制約・回答ルールに沿って内容を解析し、判断を支援する",
         "判断結果・確認観点", "MVP範囲を超えた回答をしない"],
        ["3", "ナレッジ・データ参照", "Copilot Studio → SharePoint",
         "ナレッジ／入力ファイル",
         "SharePointのナレッジ・リスト・ライブラリを参照する",
         "根拠となる資料名", "外部Web・APIは参照しない"],
        ["4", "出力生成", "Copilot Studio", "判断結果・参照結果",
         "確認可能な形式1つで出力を生成する", out,
         "根拠を必ず提示する"],
        ["5", "保存・通知", "Power Automate（必要時のみ）", "生成された出力",
         "標準コネクタでSharePoint保存・Teams通知・一覧更新を行う",
         "保存済みデータ・通知",
         "エージェントフローを呼ぶ場合は When an agent calls the flow"],
        ["6", "人の確認・修正", "人", "出力・通知",
         "内容を確認し、誤り・不足を修正して確定する",
         "確定した成果物", "自動確定せず必ず人が確認する"],
    ]


def arch_rows(tid, d):
    v = V1T[tid]
    apps = "／".join(d["overview"]["used_apps"])
    return [
        ["利用者 / メンター", "入力の投入と最終確認", "人", "―"],
        ["Copilot Studio", "意図理解・判断支援・生成・案内・確認観点の提示",
         "エージェント（%s）" % d["metadata"]["agent_name"],
         "MVPに必ず含める"],
        ["SharePoint / OneDrive", "ナレッジ・ファイル・リストの標準保管先",
         "リスト／ライブラリ", "共有ファイルサーバ直接参照は避ける"],
        ["Power Automate", "保存・通知・一覧更新（必要最小限）",
         "標準コネクタのみ",
         "エージェントフロー呼び出しは When an agent calls the flow"],
        ["出力", v["mermaid_out"], "確認可能な形式1つ", "―"],
        ["人の確認・修正", "出力の確認と修正、例外対応", "人", "必ず設ける"],
        ["使用アプリ（本チーム）", apps, "―", "詳細は 05_使用アプリ一覧"],
    ]


def build_steps_rows(tid):
    v = V1T[tid]
    ref = {"Step 1": "SPO_リスト設計／SPO_ライブラリ設計",
           "Step 2": "08_ダミーデータ一覧",
           "Step 3": "CS_CopilotStudio設計",
           "Step 4": "PA_PowerAutomate設計",
           "Step 5": "09_テストケース"}
    rows = []
    for st in v["steps"]:
        head = st["title"]
        key = head.split(":")[0].strip()
        rows.append([
            key,
            head.split(": ", 1)[1] if ": " in head else head,
            "\n".join("・" + i.replace("`", "") for i in st["items"]),
            ref.get(key, "―"),
        ])
    return rows


def role_rows(tid):
    v = V1T[tid]
    detail = {
        "人": "入力データの準備、出力の確認・承認、例外時の判断。"
              "自動確定はさせない。",
        "Copilot Studio": "利用者の入力を解釈し、ナレッジ・入力ファイルを"
                          "参照して判断と生成を行う。根拠を提示する。",
        "Power Automate": "標準コネクタでの保存・通知・一覧更新のみ。"
                          "エージェントフロー呼び出しは "
                          "When an agent calls the flow を使う。",
        "SharePoint/OneDrive/Excel": "入力データ・ナレッジ・出力結果の保管と"
                                     "整形。命名規則と権限を先に決める。",
    }
    rows = []
    for raw in v["roles"]:
        label, _, body = raw.partition(":")
        label = label.strip()
        rows.append([label, body.strip(), detail.get(label, "")])
    return rows


def unresolved_rows(tid, d):
    """12_未決事項・事務局確認。V1の横断論点＋品質レビュー＋チーム固有の確認"""
    rows = []
    for num_ttl, body in ATTENTION:
        num, ttl = num_ttl.split(". ", 1)
        target = {"1": "全チーム", "2": "T08 / T09", "3": "全チーム",
                  "4": "全チーム（特にT03）", "5": "T09",
                  "6": "T03 / T06 / T08 / T09"}.get(num, "全チーム")
        hit = "該当" if (tid in target or "全チーム" in target) else "非該当"
        rows.append(["横断論点 " + num, ttl, body, target, hit, "運営"])
    for a in QUALITY_ADD:
        owner = "事務局" if ("個人情報" in a or "メール" in a) else "運営"
        hit = "該当" if (tid in a or "各チーム" in a or "Agent Flow" in a) \
            else "参考"
        rows.append(["品質レビュー", "追加確認が望ましい点", a, "―", hit, owner])
    for r in d["risks_and_advice"]:
        rows.append(["本チーム固有", r["リスク"],
                     "%s（影響：%s）" % (r["対応案"], r["影響"]),
                     d["metadata"]["team_id"], "該当",
                     "メンター／" + r["確認タイミング"]])
    return rows


ALT_APP = {
    "Forms": "Copilot Studioの会話入力／アダプティブカードで代替",
    "Power BI": "Excel Power Query・ピボット・SharePointリストビューで代替",
    "Dataverse": "SharePointリスト／Excelで代替",
    "Premium Connector": "標準コネクタの範囲で構成する",
    "外部Web/API": "人手取得＋SharePoint格納で回避する",
    "外部API": "使用しない",
    "外部Web検索": "人手取得＋SharePoint格納で回避する",
}
APP_SHEET = {
    "Copilot Studio": "CS_CopilotStudio設計",
    "SharePointリスト": "SPO_リスト設計",
    "SharePointライブラリ": "SPO_ライブラリ設計",
    "Power Automate": "PA_PowerAutomate設計",
    "Teams": "Teams_通知設計",
    "Outlook": "Outlook_メール予定表設計",
    "Excel": "Excel_データ処理設計",
    "Word": "WordPPT_出力設計",
    "PowerPoint": "WordPPT_出力設計",
    "OneDrive": "OneDrive_保管設計",
}
UNDECIDED = "（未確定：メンター試作時に確定）"


def frame_rows(name, tid, d):
    """設計データがJSONに無いアプリ固有シートの初期枠"""
    v = V1T[tid]
    out1 = v["outputs"].split("、")[0]
    if name == "Teams_通知設計":
        return [[1, "処理結果の通知",
                 "Copilot Studioの出力を人が確認した後",
                 "担当チャネル／担当者チャット " + UNDECIDED,
                 "Teams（標準コネクタ）",
                 "%s の要点＋SharePointリンク %s" % (out1, UNDECIDED),
                 "SharePointリスト／ライブラリのリンク", "メンター",
                 "自動送信の可否は内容次第。実データ・個人情報は載せない",
                 "通知先と文面が事務局の運用ルールに沿うか"]]
    if name == "Outlook_メール予定表設計":
        return [[1, "メール文案の送信", "メール",
                 "人が内容を確認・承認した後", "宛先 " + UNDECIDED,
                 "件名テンプレート " + UNDECIDED,
                 "本文テンプレート " + UNDECIDED,
                 "人の確認後に送信（自動送信はしない）",
                 "予定表・会議室の自動参照は事務局確認後に判断する",
                 "送信可否の判断基準が明確か"],
                [2, "予定表の参照", "予定表",
                 "会議調整が必要になったとき", "―", "―", "―",
                 "人が予定表を確認して招集する",
                 "予定表・会議室の自動参照はMVP対象外（事務局確認待ち）",
                 "自動招集に寄っていないか"]]
    if name == "Excel_データ処理設計":
        return [[1, "整形・集計用ブック " + UNDECIDED, "Data",
                 "Power Query／ピボットで整形・集計する（Power BIは使わない）",
                 "入力範囲 " + UNDECIDED, "出力範囲 " + UNDECIDED,
                 "人が更新（定時自動更新はMVP対象外）",
                 "実データ・個人情報は置かない。ダミーデータで検証する",
                 "Excel／SharePointリストビューだけで代替できているか"]]
    if name == "WordPPT_出力設計":
        return [[1, out1, "テンプレート " + UNDECIDED,
                 "差し込み項目 " + UNDECIDED,
                 "Copilot Studioが要点・構成案を生成し、人が仕上げる",
                 "SharePoint／OneDrive", "YYYYMMDD_概要_版数",
                 "人が最終確認して確定する",
                 "完全自動生成は非MVP。要点・構成案までに絞る"]]
    if name == "OneDrive_保管設計":
        return [[1, "01_Input", "入力資料の集約",
                 v["inputs"], "YYYYMMDD_概要_版数", "メンター／チーム内",
                 "MVP検証期間のみ保持", "実データは置かない"],
                [2, "02_Output", "生成物の保管", v["outputs"],
                 "YYYYMMDD_概要_版数", "メンター／チーム内",
                 "MVP検証期間のみ保持", "人の確認後の版を保管する"],
                [3, "03_Knowledge", "ナレッジ・基準の保管",
                 "チェック基準・FAQ・手順書など " + UNDECIDED,
                 "YYYYMMDD_概要_版数", "メンター／チーム内",
                 "MVP検証期間のみ保持", "SharePointライブラリとの役割を分ける"]]
    return []


SHEET_WIDTHS = {
    "00_表紙・前提": [30, 34, 96],
    "01_企画概要": [26, 100, 40],
    "02_MVPスコープ": [18, 26, 100],
    "03_全体業務フロー": [6, 18, 26, 26, 46, 26, 40],
    "04_アーキテクチャ": [26, 54, 28, 42],
    "05_使用アプリ一覧": [22, 12, 24, 42, 26, 44],
    "06_人AI自動化分担": [26, 46, 70],
    "07_構築ステップ総覧": [10, 24, 76, 34],
    "08_ダミーデータ一覧": [12, 30, 42, 12, 14, 14, 14, 22],
    "09_テストケース": [8, 12, 34, 44, 32],
    "10_リスク・躓きポイント": [28, 10, 62, 24],
    "11_メンター確認チェックリスト": [6, 70, 10, 32],
    "12_未決事項・事務局確認": [16, 28, 70, 20, 10, 26],
    "13_変更履歴": [10, 14, 80, 26],
    "CS_CopilotStudio設計": [12, 22, 60, 50, 32, 32],
    "SPO_リスト設計": [18, 30, 18, 18, 14, 8, 22, 12, 10, 10, 30, 22, 14, 20],
    "SPO_ライブラリ設計": [22, 34, 16, 30, 22, 20, 26],
    "PA_PowerAutomate設計": [6, 18, 36, 30, 28, 20, 30, 24, 24, 22, 26, 28,
                            34, 30],
}
for _k, (_c, _w) in FRAME_COLS.items():
    SHEET_WIDTHS[_k] = _w


def build_team(tid, images):
    d = json.load(open(os.path.join(IN, "data", "%s_design_data.json" % tid),
                       encoding="utf-8"))
    v = V1T[tid]
    m = d["metadata"]
    ctx = {"team_id": m["team_id"], "team_name": m["team_name"],
           "agent_name": m["agent_name"], "design_status": m["design_status"]}
    app_sheets = list(d["app_specific_sheets"])
    sheets = COMMON_SHEETS + app_sheets

    # 画像をシート別に割り当てる
    by_sheet = {}
    for vtype, path in images[tid]:
        target = IMG_SHEET.get(vtype)
        if target is None or target not in sheets:
            target = "04_アーキテクチャ"      # 想定外は必ずどこかに載せる
        by_sheet.setdefault(target, []).append((vtype, path))

    wb = Workbook()
    wb.remove(wb.active)

    def new(name):
        return Sheet(wb, name, ctx, SHEET_WIDTHS[name])

    def imgs(s):
        for vtype, path in by_sheet.get(s.name, []):
            s.image(path, IMG_CAPTION.get(vtype, vtype), "")

    banner = os.path.join(IN, d["overview"]["banner_image"])

    # ---------------- 00_表紙・前提（ダッシュボード型） ----------------
    s = new("00_表紙・前提")
    s.section("本設計書の対象と前提")
    s.table(["区分", "項目", "内容"], [
        ["対象", "チームID / チーム名", "%s / %s" % (m["team_id"], m["team_name"])],
        ["対象", "エージェント名", m["agent_name"]],
        ["対象", "類型", m["pattern"]],
        ["対象", "設計ステータス", m["design_status"]],
        ["対象", "対象フェーズ", m["対象フェーズ"]],
        ["対象", "版 / 作成日 / 作成", "%s / %s / %s" % (VERSION, MADE_ON, VENDOR)],
        ["前提", "目的", v["purpose"]],
        ["前提", "MVPゴール", d["overview"]["mvp"]],
        ["前提", "使用アプリ", "／".join(d["overview"]["used_apps"])],
        ["前提", "読者", "メンター（Power Platformの基本的な知識がある前提）"],
        ["前提", "正本", "priority_medium_excel_design_input_v1 の設計データ、"
                       "および v3 高品質ビジュアル"],
    ], ink_cols=("内容",), filt=False)

    s.image(banner, "%s 全体アーキテクチャ（V1バナー）" % tid, "")

    s.section("共通制約（全チーム共通・逸脱していないか必ず確認する）", "blue")
    s.table(["区分", "項目", "内容"],
            [["共通制約", "No.%d" % (i + 1), c]
             for i, c in enumerate(CONSTRAINTS)], filt=False)
    s.note("Power Automate 側でエージェントフローを呼び出す場合は "
           "When an agent calls the flow を使用する。"
           "通常トリガーからの直接呼び出しは行わない。", "warn")
    s.blank()

    s.section("共通アーキテクチャ原則", "blue")
    s.table(["区分", "項目", "内容"],
            [["原則", "No.%d" % (i + 1), p] for i, p in enumerate(PRINCIPLES)],
            filt=False)

    s.section("収録シート一覧", "blue")
    s.table(["区分", "項目", "内容"],
            [["共通シート" if n in COMMON_SHEETS else "アプリ固有シート",
              n, SHEET_PURPOSE.get(n, "")] for n in sheets], filt=False)

    s.section("T04の扱い（本設計書の対象外）", "risk")
    s.table(["区分", "項目", "内容"],
            [["対象外", "T04", T04_NOTE[0]],
             ["対象外", "今後の代替案", T04_NOTE[1] if len(T04_NOTE) > 1 else ""]],
            filt=False)
    s.finish()

    # ---------------- 01_企画概要 ----------------
    s = new("01_企画概要")
    s.section("企画概要")
    s.table(["項目", "内容", "備考"], [
        ["類型", m["pattern"], "V1の類型と同一"],
        ["目的", v["purpose"], ""],
        ["概要", d["overview"]["summary"], ""],
        ["MVPの考え方", d["overview"]["mvp"], "詳細は 02_MVPスコープ"],
        ["入力", v["inputs"], "ダミーデータで代替（08_ダミーデータ一覧）"],
        ["出力", v["outputs"], "確認可能な形式1つに絞る"],
        ["使用アプリ", "／".join(d["overview"]["used_apps"]),
         "詳細は 05_使用アプリ一覧"],
        ["制約への適合", v["fit"], "共通制約は 00_表紙・前提"],
    ], ink_cols=("内容",), filt=False)
    imgs(s)
    s.finish()

    # ---------------- 02_MVPスコープ ----------------
    s = new("02_MVPスコープ")
    s.section("MVPスコープ")
    rows = [["MVPゴール", "ゴール", d["overview"]["mvp"]]]
    rows += [["MVPスコープ", "含める", x] for x in v["mvp_scope"]]
    rows += [["非MVP範囲", "含めない", x] for x in v["mvp_out"]]
    rows += [["前提", "対象フェーズ", m["対象フェーズ"]],
             ["前提", "完成度", "完成版ではなく、一本通して動くことを優先する"]]
    s.table(["区分", "項目", "内容"], rows, ink_cols=("内容",), filt=False)
    s.note("非MVP範囲に手を広げると8月中に一本通らなくなる。"
           "1ユースケースに絞ること。", "warn")
    imgs(s)
    s.finish()

    # ---------------- 03_全体業務フロー ----------------
    s = new("03_全体業務フロー")
    s.section("全体業務フロー（入力 → AI処理 → 保管 → 出力 → 人の確認）")
    fr = flow_rows(tid, d)
    fr[0][3] = v["inputs"]
    s.table(["No", "工程", "主体", "入力", "処理内容", "出力", "確認ポイント・注意"],
            fr, height=44)
    s.note("本シートの工程は、V1のMermaidアーキテクチャと人/AI/自動化の分担から"
           "起こしたもの。図と矛盾しないこと。")
    imgs(s)
    s.finish()

    # ---------------- 04_アーキテクチャ ----------------
    s = new("04_アーキテクチャ")
    s.section("全体アーキテクチャ図")
    s.image(banner, "%s 全体アーキテクチャ（V1バナー）" % tid, "")
    s.section("構成要素と役割")
    s.table(["構成要素", "役割", "実体", "備考"], arch_rows(tid, d), height=34)
    s.section("Mermaid定義（V1原文）", "blue")
    for ln in v["mermaid"].splitlines():
        s.note("　" + ln if ln.startswith(" ") else ln)
    imgs(s)
    s.finish()

    # ---------------- 05_使用アプリ一覧 ----------------
    s = new("05_使用アプリ一覧")
    s.section("使用アプリと、使用しないアプリの理由")
    rows = []
    for a in d["app_usage"]:
        app = a["アプリ"]
        sh = APP_SHEET.get(app)
        if a["利用有無"] == "使用":
            note = ("作成： " + sh) if sh in app_sheets else \
                "固有シートなし（本シートに用途を記載）"
        else:
            note = "作成しない（未使用）"
        rows.append([app, a["利用有無"], a["用途"], a["不要理由"], note,
                     ALT_APP.get(app, "")])
    s.table(["アプリ", "利用有無", "用途", "不要理由", "アプリ固有シート",
             "代替手段"], rows, height=32)
    mapped = {APP_SHEET.get(a["アプリ"]) for a in d["app_usage"]
              if a["利用有無"] == "使用"}
    orphan = [x for x in app_sheets if x not in mapped]
    if orphan:
        s.blank()
        s.section("使用アプリ一覧に対応行が無いアプリ固有シート（突合用）", "blue")
        s.table(["アプリ", "利用有無", "用途", "不要理由", "アプリ固有シート",
                 "代替手段"],
                [[{"SPO_リスト設計": "SharePointリスト",
                   "SPO_ライブラリ設計": "SharePointライブラリ"}.get(x, "―"),
                  "使用", "本入力JSONの app_specific_sheets に含まれるため作成",
                  "―", x,
                  "―"] for x in orphan], filt=False)
        s.note("入力JSONの app_usage 側に行が無いが、app_specific_sheets に"
               "含まれるため作成している。使用有無をメンター試作時に確認すること。",
               "warn")
    s.blank()
    s.note(d["unused_app_sheet_policy"])
    s.note("Forms / Power BI / Dataverse / Premium Connector / 外部Web・API は"
           "使用しない。代替手段の列を必ず確認すること。", "warn")
    miss = [x for x in app_sheets
            if x in FRAME_COLS or (x == "PA_PowerAutomate設計"
                                   and not d["power_automate_design"])
            or (x == "SPO_ライブラリ設計"
                and not d["sharepoint_library_design"])]
    if miss:
        s.blank()
        s.section("設計データが未整備のアプリ固有シート", "risk")
        s.table(["アプリ", "利用有無", "用途", "不要理由", "アプリ固有シート",
                 "代替手段"],
                [["―", "使用", "本入力JSONに設計行が無いため設計枠のみ作成",
                  "―", x, "メンター試作時に確定する"] for x in miss],
                filt=False)
    s.finish()

    # ---------------- 06_人AI自動化分担 ----------------
    s = new("06_人AI自動化分担")
    s.section("人 / AI / 自動化の分担")
    s.table(["主体", "担当範囲（V1原文）", "具体的な作業と注意"],
            role_rows(tid), height=48)
    s.note("人の確認ポイントを必ず設ける。AIの出力をそのまま確定させない。", "warn")
    imgs(s)
    s.finish()

    # ---------------- 07_構築ステップ総覧 ----------------
    s = new("07_構築ステップ総覧")
    s.section("メンター試作 構築ステップ（全チーム共通・V1原文）")
    s.table(["STEP", "作業", "作業内容", "参照シート"],
            build_steps_rows(tid), height=56)
    s.blank()
    s.section("本チームで先に決めること", "blue")
    s.table(["STEP", "作業", "作業内容", "参照シート"], [
        ["事前", "格納場所", "SharePointのライブラリ／リスト、命名規則、権限を"
                          "先に決める", "SPO_リスト設計／SPO_ライブラリ設計"],
        ["事前", "入力と出力", "入力： %s\n出力： %s" % (v["inputs"], v["outputs"]),
         "01_企画概要"],
        ["事前", "ダミーデータ", "実データ・個人情報を使わないダミーを用意する",
         "08_ダミーデータ一覧"],
    ], height=48, filt=False)
    imgs(s)
    s.finish()

    # ---------------- 08_ダミーデータ一覧 ----------------
    s = new("08_ダミーデータ一覧")
    s.section("ダミーデータ（サンプル行）")
    cols = list(d["dummy_data"][0].keys())
    s.table(cols, [[r.get(c, "") for c in cols] for r in d["dummy_data"]],
            ink_cols=tuple(cols), height=30)
    s.section("サンプルファイルとデータ方針", "blue")
    for f in v["dummy_files"]:
        s.note("サンプルファイル： " + f.replace("`", ""))
    for p in SAMPLE_POLICY:
        s.note("方針： " + p)
    s.note("推奨列： " + "／".join(v["dummy_cols"]))
    s.note("実データ・個人情報・実顧客名は使用しない。カテゴリ名だけで表現する。",
           "warn")
    imgs(s)
    s.finish()

    # ---------------- 09_テストケース ----------------
    s = new("09_テストケース")
    s.section("最低限テストケース（基本ケース中心）")
    cols = list(d["test_cases"][0].keys())
    s.table(cols, [[r.get(c, "") for c in cols] for r in d["test_cases"]],
            height=34)
    s.note("AI出力は必ず人が確認し、誤りや不足を記録する。"
           "外部Web・APIを使わずに実行できることも確認する。", "warn")
    imgs(s)
    s.finish()

    # ---------------- 10_リスク・躓きポイント ----------------
    s = new("10_リスク・躓きポイント")
    s.section("リスクと対応案")
    cols = list(d["risks_and_advice"][0].keys())
    s.table(cols, [[r.get(c, "") for c in cols] for r in d["risks_and_advice"]],
            height=40)
    s.blank()
    s.section("躓きポイントとメンターからの助言（V1原文）", "risk")
    rows = [[r, "―", "―", "―"] for r in v["risks"]]
    for i, a in enumerate(v["advice"]):
        if i < len(rows):
            rows[i][2] = a
    s.table(["躓きポイント", "影響", "メンターからの助言", "確認タイミング"],
            [[r[0], "―", r[2], "初回試作時"] for r in rows], height=32,
            filt=False)
    imgs(s)
    s.finish()

    # ---------------- 11_メンター確認チェックリスト ----------------
    s = new("11_メンター確認チェックリスト")
    s.section("本チームの確認項目")
    cols = list(d["mentor_checklist"][0].keys())
    s.table(cols, [[r.get(c, "") for c in cols] for r in d["mentor_checklist"]],
            height=24)
    s.blank()
    s.section("試作前に確認する（制約の適合）", "blue")
    s.table(["No", "確認項目", "完了", "備考"], [
        [1, "Copilot StudioをMVPに含めているか", "", ""],
        [2, "Forms／Power BI／Dataverse／Premium Connector／外部Web・APIを"
            "使っていないか", "", "05_使用アプリ一覧"],
        [3, "SharePoint／OneDrive格納になっているか（共有ファイルサーバ"
            "直接参照でないか）", "", ""],
        [4, "Power Automateからエージェントフローを呼ぶ場合、"
            "When an agent calls the flow になっているか", "",
         "PA_PowerAutomate設計"],
        [5, "ダミーデータに実名・実データ・個人情報が含まれていないか", "",
         "08_ダミーデータ一覧"],
    ], height=24, filt=False)
    s.blank()
    s.section("試作後に確認する（MVPの成立・V1原文）", "blue")
    s.table(["No", "確認項目", "完了", "備考"],
            [[i + 1, c, "", ""] for i, c in enumerate(CHECKLIST)],
            height=22, filt=False)
    s.blank()
    s.section("先回り検証ポイント（V1原文）", "blue")
    s.table(["No", "確認項目", "完了", "備考"],
            [[i + 1, c, "", ""] for i, c in enumerate(v["precheck"])],
            height=22, filt=False)
    imgs(s)
    s.finish()

    # ---------------- 12_未決事項・事務局確認 ----------------
    s = new("12_未決事項・事務局確認")
    s.section("未決事項と確認先")
    s.table(["区分", "論点", "内容", "対象", "本チームの該当", "確認先"],
            unresolved_rows(tid, d), height=40)
    s.note("「該当」の行は、メンター試作の前または初回試作時に確認する。", "warn")
    imgs(s)
    s.finish()

    # ---------------- 13_変更履歴 ----------------
    s = new("13_変更履歴")
    s.section("変更履歴")
    s.table(["版", "日付", "変更内容", "作成"], [
        [VERSION, MADE_ON,
         "初版作成。priority_medium_excel_design_input_v1 の設計データと "
         "v3 高品質ビジュアル、およびメンター向け設計資料一式（V1）をもとに、"
         "共通14シート＋アプリ固有シートを作成。", VENDOR],
    ], height=48, filt=False)
    s.note("本設計書はメンター試作用。実装確定後はチーム側で更新すること。")
    s.finish()

    # ---------------- アプリ固有シート ----------------
    for name in app_sheets:
        s = new(name)
        if name == "CS_CopilotStudio設計":
            s.section("Copilot Studio エージェント設計")
            cols = list(d["copilot_studio_design"][0].keys())
            s.table(cols, [[r.get(c, "") for c in cols]
                           for r in d["copilot_studio_design"]], height=44)
            s.blank()
            s.section("システム指示に必ず入れる制約", "blue")
            for c in ["外部Web・API・スクレイピングは参照しない",
                      "参照先はSharePointのナレッジ・リスト・ライブラリに限定する",
                      "判断が分かれる場合は「要確認」として人に返す",
                      "回答には根拠となる資料名を必ず添える",
                      "MVP範囲を超えた回答をしない"]:
                s.note("・" + c, "ink")
            s.note("Formsは使用しない。入力は会話入力またはアダプティブカードで"
                   "受け取る。", "warn")
        elif name == "SPO_リスト設計":
            s.section("SharePointリスト 列定義")
            if d["sharepoint_list_design"]:
                cols = list(d["sharepoint_list_design"][0].keys())
                s.table(cols, [[r.get(c, "") for c in cols]
                               for r in d["sharepoint_list_design"]],
                        ink_cols=("入力例", "既定値", "選択肢"), height=30)
            else:
                s.note("本入力JSONに列定義がありません。" + UNDECIDED, "warn")
            s.note("内部名は作成後に変更できない。列表示名と内部名を分けて管理する。",
                   "warn")
        elif name == "SPO_ライブラリ設計":
            s.section("SharePointライブラリ設計")
            if d["sharepoint_library_design"]:
                cols = list(d["sharepoint_library_design"][0].keys())
                s.table(cols, [[r.get(c, "") for c in cols]
                               for r in d["sharepoint_library_design"]],
                        ink_cols=("命名規則",), height=30)
            else:
                s.table(["ライブラリ名", "用途", "フォルダ案", "格納物",
                         "命名規則", "権限", "備考"],
                        [["ライブラリ名 " + UNDECIDED, v["inputs"], "01_Input",
                          "入力ファイル・サンプルデータ", "YYYYMMDD_概要_版数",
                          "メンター編集可", "実データは置かない"]],
                        height=30, filt=False)
                s.note("本入力JSONに設計行が無いため設計枠のみ。"
                       "メンター試作時に確定する。", "warn")
        elif name == "PA_PowerAutomate設計":
            s.section("Power Automate 処理設計（B粒度：実装前設計レベル）")
            if d["power_automate_design"]:
                cols = list(d["power_automate_design"][0].keys())
                s.table(cols, [[r.get(c, "") for c in cols]
                               for r in d["power_automate_design"]], height=48)
            else:
                s.note("本入力JSONに処理設計が無いため設計枠のみ。" + UNDECIDED,
                       "warn")
            s.note("エージェントフローを呼び出す場合のトリガーは "
                   "When an agent calls the flow のみ。"
                   "通常トリガーから直接呼び出すと従量課金の懸念があるため、"
                   "SPOリスト作成・フォルダへのダミーファイル作成・人の確認操作"
                   "などの間接起動で回避できるか検討する"
                   "（従量課金を回避できることが確認できる場合に限る）。", "warn")
            s.note("プレミアムコネクタは使用しない。標準コネクタ"
                   "（SharePoint／Teams／Outlook／Excel）のみで構成する。", "warn")
        else:
            cols, _w = FRAME_COLS[name]
            s.section(SHEET_PURPOSE.get(name, name) + "（設計枠）")
            s.table(cols, frame_rows(name, tid, d), height=44)
            s.note("本入力JSONに設計行が無いため、列枠と初期案のみを記載している。"
                   "「%s」の箇所はメンター試作時に確定すること。" % UNDECIDED,
                   "warn")
        imgs(s)
        s.finish()

    os.makedirs(OUTDIR, exist_ok=True)
    path = os.path.join(OUTDIR, "%s_detailed_design.xlsx" % tid)
    wb.save(path)
    return path, sheets, {k: len(vv) for k, vv in by_sheet.items()}


def main():
    images = load_images()
    total = 0
    for tid in TEAMS:
        p, sheets, place = build_team(tid, images)
        mb = os.path.getsize(p) / 1e6
        total += mb
        print("%s: %2dシート / 追加画像%2d枚 / %.1f MB" % (
            tid, len(sheets), sum(place.values()), mb))
    print("合計 %.0f MB" % total)


if __name__ == "__main__":
    main()
