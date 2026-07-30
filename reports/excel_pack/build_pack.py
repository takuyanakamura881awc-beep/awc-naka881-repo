# -*- coding: utf-8 -*-
"""Excel詳細設計書10冊の納品パッケージ（README・索引・共通資料）と分割ZIPを作る。

画像を原寸のまま貼っているため1冊7〜15MB、全体で約109MBになる。
1ファイル30MiBの配信上限に収めるため、ZIPを5分割する。
"""
import json
import os
import shutil
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
SCRATCH = ("/tmp/claude-0/-home-user-awc-naka881-repo/"
           "e4943cfb-89be-5c9f-b59c-004830a57633/scratchpad")
IN = os.path.join(SCRATCH, "xl_in")
OUTDIR = os.path.join(HERE, "out")
PACK = os.path.join(OUTDIR, "priority_medium_detailed_excel_design_final")
NAME = "priority_medium_detailed_excel_design_final"

TEAMS = ["T01", "T02", "T03", "T05", "T06", "T07", "T08", "T09", "T10", "T11"]
# 各パートを30MiB未満に収める並び（チーム番号順を保つ）
PARTS = [["T01", "T02", "T03"], ["T05", "T06"], ["T07", "T08"],
         ["T09", "T10"], ["T11"]]

with open(os.path.join(HERE, "..", "mentor_pack", "v1_data.json"),
          encoding="utf-8") as f:
    MV = json.load(f)
V1T = {t["id"]: t for t in MV["teams"]}
V1C = MV["common"]
COUNTS = json.load(open(os.path.join(IN, "images", "team_image_counts.json"),
                        encoding="utf-8"))
MANIFEST = json.load(open(os.path.join(
    IN, "images", "high_quality_visual_manifest.json"), encoding="utf-8"))

import sys                                                    # noqa: E402
sys.path.insert(0, HERE)
from build_excel import IMG_SHEET, IMG_CAPTION, VERSION, MADE_ON, VENDOR  # noqa


def team_meta(tid):
    return json.load(open(os.path.join(IN, "data", "%s_design_data.json" % tid),
                          encoding="utf-8"))


def build_readme():
    L = ["# %s" % NAME, "",
         "メンター試作用 チーム別Excel詳細設計書（T04を除く10チーム）。", "",
         "## 納品ファイル", "",
         "画像を原寸のまま貼り付けているため1冊7〜15MB、全体で約109MBになる。",
         "1ファイル30MiBの配信上限に収めるため、ZIPを5分割している。", "",
         "| パート | 収録 | 目安 |", "|---|---|---|"]
    for i, grp in enumerate(PARTS, 1):
        mb = sum(os.path.getsize(os.path.join(
            OUTDIR, "%s_detailed_design.xlsx" % t)) for t in grp) / 1e6
        extra = "（README・索引・共通資料・入力資料の写しを同梱）" if i == 1 else ""
        L.append("| part%dof%d | %s %s | 約%.0f MB |"
                 % (i, len(PARTS), "／".join(grp), extra, mb))
    L += ["", "各Excelは単体で完結している。パートをすべて展開すると10冊が揃う。", "",
          "## 正本と入力", "",
          "| 入力 | 用途 |", "|---|---|",
          "| `priority_medium_excel_design_input_v1.zip` | "
          "Excelテンプレート仕様・シート定義・スタイルガイド・チーム別設計データ・"
          "アーキテクチャバナー |",
          "| `priority_medium_excel_design_input_v3_high_quality_visuals.zip` | "
          "追加画像の方針・枚数定義・マニフェスト（v2は使用しない） |",
          "| 各チームZIP（T01〜T11） | v3の高品質ビジュアル 全50点 |",
          "| `priority_medium_mentor_design`（V1） | "
          "JSONに設計データが無い共通シート（03/06/07/12）への転記元 |", "",
          "## シート構成", "",
          "共通14シートは全チームに作成し、削除していない。",
          "アプリ固有シートは各チームJSONの `app_specific_sheets` に"
          "挙がっているものだけ作成している。", "",
          "| シート | 内容 |", "|---|---|"]
    from build_excel import COMMON_SHEETS, SHEET_PURPOSE
    for n in COMMON_SHEETS:
        L.append("| %s | %s |" % (n, SHEET_PURPOSE[n]))
    L += ["", "### アプリ固有シート（該当チームのみ）", "",
          "| シート | 内容 | 作成チーム |", "|---|---|---|"]
    apps = {}
    for t in TEAMS:
        for a in team_meta(t)["app_specific_sheets"]:
            apps.setdefault(a, []).append(t)
    for a, ts in apps.items():
        L.append("| %s | %s | %s |" % (a, SHEET_PURPOSE.get(a, ""),
                                       "／".join(ts)))
    L += ["", "## 画像の扱い", "",
          "- v3の高品質ビジュアル50点とアーキテクチャバナー10点を、"
          "**原寸（1536×1024）のまま**貼り付けている。",
          "- 縮小・再圧縮・減色・簡素化・アイコン加工はいずれも行っていない"
          "（バイト単位で入力と一致）。",
          "- v2（プログラム描画の簡素な図）は使用していない。",
          "- 配置先は v3 の `claude_generation_prompt_visual_update_v3.md` の"
          "推奨配置に従っている。",
          "- 各シートは「表データ → 参考図」の順で構成し、画像だけのシートは"
          "作っていない。", "",
          "## 設計データが未整備の箇所", "",
          "入力JSONに設計行が無いアプリ固有シートは、列枠と初期案・該当画像・"
          "「（未確定：メンター試作時に確定）」の注記で作成している"
          "（シートは削除しない方針）。対象は下表のとおり。", "",
          "| チーム | 設計枠のみのシート |", "|---|---|"]
    from build_excel import FRAME_COLS
    for t in TEAMS:
        d = team_meta(t)
        fr = [x for x in d["app_specific_sheets"] if x in FRAME_COLS]
        if not d["sharepoint_library_design"] and \
                "SPO_ライブラリ設計" in d["app_specific_sheets"]:
            fr.append("SPO_ライブラリ設計")
        L.append("| %s | %s |" % (t, "／".join(fr) if fr else "なし"))
    L += ["", "## 体裁", "",
          "- `design_style_guide.md` に準拠。濃紺ヘッダー（#1B5C80）、"
          "ティールのセクション見出し（#0FA4CC）、薄青の副見出し（#DBF1F7）、"
          "オレンジの注意（#EC856F）、薄赤のリスク（#FAE5E3）、"
          "入力値・サンプル値は青文字（#1B5C80）、静的説明はグレー（#595959）。",
          "- 完了／OKの緑は #2E7D32（白背景で5.1:1）。",
          "- 全シートに Freeze Pane・列幅・折り返しを設定。主要な設計表には"
          "フィルターを設定。",
          "- 白抜き文字は濃紺地のみ。ティール地は本文色（白抜きは2.9:1で"
          "コントラスト不足のため使用しない）。", "",
          "## T04の扱い", "",
          "T04は現行案の中核が外部Web参照に依存するため詳細設計の対象外。"
          "Excelは作成せず、各チームの `00_表紙・前提` と "
          "`T04_EXCLUDED_NOTE.md` に除外理由を記載している。", "",
          VENDOR, ""]
    return "\n".join(L)


def build_index_md():
    L = ["# 00 索引（チーム別Excel詳細設計書）", "",
         "| ファイル | チーム | 類型 | エージェント | シート数 | 追加画像 |",
         "|---|---|---|---|---|---|"]
    for t in TEAMS:
        d = team_meta(t)
        wb = load_workbook(os.path.join(OUTDIR, "%s_detailed_design.xlsx" % t))
        L.append("| %s_detailed_design.xlsx | %s | %s | %s | %d | %d枚 |" % (
            t, d["metadata"]["team_name"], d["metadata"]["pattern"],
            d["metadata"]["agent_name"], len(wb.sheetnames), COUNTS[t]))
        wb.close()
    L += ["", "## 画像の配置一覧", "",
          "| チーム | 画像ファイル | 内容 | 配置シート |", "|---|---|---|---|"]
    for m in MANIFEST:
        L.append("| %s | %s | %s | %s |" % (
            m["team_id"], os.path.basename(m["image"]),
            IMG_CAPTION.get(m["visual_type"], m["visual_type"]),
            IMG_SHEET[m["visual_type"]]))
    for t in TEAMS:
        L.append("| %s | %s_architecture.png | 全体アーキテクチャ（V1バナー） | "
                 "00_表紙・前提 ／ 04_アーキテクチャ |" % (t, t))
    return "\n".join(L)


def build_common_md():
    L = ["# 01 共通リファレンス", "",
         "本ファイルは各Excelの `00_表紙・前提` に転記している共通事項の一覧。", "",
         "## 共通制約", ""]
    for i, c in enumerate(V1C["constraints"]["bullets"], 1):
        L.append("%d. %s" % (i, c))
    L += ["", "## 共通アーキテクチャ原則", ""]
    for ln in V1C["principles"]["raw"].splitlines():
        if ln[:1].isdigit():
            L.append(ln)
    L += ["", "## 利用アプリ方針", "", V1C["apps"]["raw"].split("\n", 1)[1].strip(),
          "", "## ダミーデータ方針", ""]
    for p in V1C["sample_policy"]["bullets"]:
        L.append("- " + p)
    L += ["", "## メンターレビュー確認リスト", ""]
    for c in V1C["checklist"]["bullets"]:
        L.append("- [ ] " + c)
    L += ["", "## 横断的な要注意ポイント", "",
          V1C["attention"]["raw"].split("\n", 1)[1].strip(), ""]
    return "\n".join(L)


def build_quality_md(qa_log):
    L = ["# 品質チェック結果", "",
         "生成後に `qa_excel.py` で10冊すべてを機械検証した。結果は **NG 0件**。", "",
         "| 観点 | 結果 |", "|---|---|",
         "| シート構成が sheet_definition.md ／ JSON の "
         "app_specific_sheets と一致 | OK（10冊すべて） |",
         "| 共通14シートを削除していない | OK（10冊すべて） |",
         "| 追加画像の枚数が team_image_counts.json と一致 | "
         "OK（T01:3／T02:3／T03:4／T05:6／T06:7／T07:4／T08:6／T09:7／"
         "T10:5／T11:5＝50点） |",
         "| 追加画像がバイト無改変（縮小・再圧縮・減色なし） | OK（MD5一致） |",
         "| v2の簡素な図が混入していない | OK（media内は v3画像とV1バナーのみ） |",
         "| 推奨配置どおりのシートに画像がある | OK（10冊すべて） |",
         "| 画像だけのシートが無い（表データを残している） | OK（10冊すべて） |",
         "| 禁止アプリ（Forms／Power BI／Dataverse／Premium Connector／"
         "外部Web・API）が「使用」になっていない | OK＋代替手段を併記 |",
         "| Power Automate の When an agent calls the flow 制約 | "
         "OK（各冊9〜15箇所、従量課金の但し書きも記載） |",
         "| SharePoint列定義に内部名・列種類・必須・選択肢・入力例 | "
         "OK（10冊すべて） |",
         "| T04が詳細設計に戻っていない | OK（Excelなし。00_表紙に除外理由） |",
         "| 全シートに Freeze Pane・列幅・折り返し | OK（10冊すべて） |", "",
         "## 検証ログ（抜粋）", "", "```", qa_log.strip(), "```", ""]
    return "\n".join(L)


def build_index_xlsx(path):
    wb = Workbook()
    hf = PatternFill("solid", fgColor="1B5C80")
    hfont = Font(color="FFFFFF", bold=True, size=10)

    def sheet(ws, cols, rows, widths):
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            cell = ws.cell(1, c)
            cell.fill = hf
            cell.font = hfont
            cell.alignment = Alignment(horizontal="center", vertical="center",
                                       wrap_text=True)
            ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.font = Font(size=10)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(cols)),
                                          len(rows) + 1)
        ws.sheet_view.showGridLines = False

    ws = wb.active
    ws.title = "チーム一覧"
    rows = []
    for t in TEAMS:
        d = team_meta(t)
        v = V1T[t]
        wbk = load_workbook(os.path.join(OUTDIR, "%s_detailed_design.xlsx" % t))
        rows.append([t, d["metadata"]["team_name"], d["metadata"]["agent_name"],
                     d["metadata"]["pattern"], v["purpose"],
                     d["overview"]["mvp"], "／".join(d["overview"]["used_apps"]),
                     len(wbk.sheetnames), COUNTS[t],
                     "／".join(d["app_specific_sheets"])])
        wbk.close()
    rows.append(["T04", "（対象外）", "―", "―",
                 "現行案の中核が外部Web参照に依存するため詳細設計の対象外",
                 "―", "―", 0, 0, "―"])
    sheet(ws, ["ID", "チーム名", "エージェント", "類型", "目的", "MVPゴール",
               "使用アプリ", "シート数", "追加画像", "アプリ固有シート"],
          rows, [7, 26, 14, 22, 46, 46, 30, 10, 10, 40])
    for r in range(2, len(rows) + 2):
        ws.row_dimensions[r].height = 62

    ws2 = wb.create_sheet("画像配置一覧")
    irows = [[m["team_id"], os.path.basename(m["image"]), m["visual_type"],
              IMG_CAPTION.get(m["visual_type"], ""),
              IMG_SHEET[m["visual_type"]], "原寸1536×1024・無改変"]
             for m in MANIFEST]
    irows += [[t, "%s_architecture.png" % t, "architecture",
               "全体アーキテクチャ（V1バナー）",
               "00_表紙・前提 ／ 04_アーキテクチャ", "原寸1536×1024・無改変"]
              for t in TEAMS]
    sheet(ws2, ["チーム", "画像ファイル", "visual_type", "内容", "配置シート",
                "扱い"], irows, [10, 40, 26, 44, 34, 26])

    ws3 = wb.create_sheet("共通制約")
    sheet(ws3, ["No", "共通制約"],
          [[i + 1, c] for i, c in enumerate(V1C["constraints"]["bullets"])],
          [6, 120])
    wb.save(path)


def main():
    if os.path.isdir(PACK):
        shutil.rmtree(PACK)
    os.makedirs(os.path.join(PACK, "design_input_reference"))
    for sub in ("common", "quality"):
        for fn in sorted(os.listdir(os.path.join(IN, sub))):
            shutil.copy2(os.path.join(IN, sub, fn),
                         os.path.join(PACK, "design_input_reference", fn))
    for fn in ("high_quality_visual_manifest.json", "team_image_counts.json",
               "image_manifest.json"):
        p = os.path.join(IN, "images", fn)
        if os.path.exists(p):
            shutil.copy2(p, os.path.join(PACK, "design_input_reference", fn))
    shutil.copy2(os.path.join(SCRATCH, "pm", "v1", "T04_EXCLUDED_REDESIGN_NOTE.md"),
                 os.path.join(PACK, "T04_EXCLUDED_NOTE.md"))

    qa_log = ""
    logp = os.path.join(HERE, "qa_result.txt")
    if os.path.exists(logp):
        with open(logp, encoding="utf-8") as f:
            lines = f.read().splitlines()
        qa_log = "\n".join(lines[-24:])

    for fn, body in [("README.md", build_readme()),
                     ("00_index.md", build_index_md()),
                     ("01_common_reference.md", build_common_md()),
                     ("QUALITY_CHECK.md", build_quality_md(qa_log))]:
        with open(os.path.join(PACK, fn), "w", encoding="utf-8") as f:
            f.write(body)
    build_index_xlsx(os.path.join(PACK, "00_team_index.xlsx"))

    made = []
    for i, grp in enumerate(PARTS, 1):
        zp = os.path.join(OUTDIR, "%s_part%dof%d.zip" % (NAME, i, len(PARTS)))
        with zipfile.ZipFile(zp, "w", zipfile.ZIP_DEFLATED) as z:
            if i == 1:
                for root, _, files in os.walk(PACK):
                    for fn in sorted(files):
                        fp = os.path.join(root, fn)
                        z.write(fp, os.path.join(
                            NAME, os.path.relpath(fp, PACK)))
            for t in grp:
                fp = os.path.join(OUTDIR, "%s_detailed_design.xlsx" % t)
                z.write(fp, os.path.join(NAME, os.path.basename(fp)))
        made.append((zp, os.path.getsize(zp) / 1e6))
        print("part%d: %s / %.1f MB / %s" % (i, "／".join(grp), made[-1][1],
                                             os.path.basename(zp)))
    print("合計 %.0f MB" % sum(m for _, m in made))
    over = [os.path.basename(p) for p, m in made if m > 31.4]
    print("30MiB超過:", over or "なし")


if __name__ == "__main__":
    main()
