# -*- coding: utf-8 -*-
"""LibreOfficeが使えないため、図形座標から簡易プレビュー画像を作って版面を確認する。
（フォント・行送りは概算。重なりや余白の確認用）"""
import io
import sys

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

SCALE = 110  # px / inch
FONTS = ["/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
         "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf",
         "/usr/share/fonts/opentype/ipafont-gothic/ipagp.ttf",
         "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"]
_F = {}


def font(pt):
    key = max(6, int(pt * SCALE / 72))
    if key not in _F:
        for p in FONTS:
            try:
                _F[key] = ImageFont.truetype(p, key)
                break
            except Exception:
                continue
        else:
            _F[key] = ImageFont.load_default()
    return _F[key]


def px(v):
    return int(Emu(v).inches * SCALE)


def fill_of(sh):
    try:
        if sh.fill.type == 1:
            return "#%06X" % int(str(sh.fill.fore_color.rgb), 16)
    except Exception:
        pass
    return None


def draw_shape(_d, im, sh):
    d = ImageDraw.Draw(im)
    if sh.left is None:
        return
    x, y = px(sh.left), px(sh.top)
    w, h = px(sh.width), px(sh.height)
    if sh.shape_type == 13:                      # 画像
        try:
            pic = Image.open(io.BytesIO(sh.image.blob)).convert("RGBA")
            pic = pic.resize((max(w, 1), max(h, 1)))
            im.paste(pic, (x, y), pic)
        except Exception:
            d.rectangle([x, y, x + w, y + h], outline="#999")
        return
    if sh.shape_type == 6:                       # グループ
        for c in sh.shapes:
            draw_shape(None, im, c)
        return
    f = fill_of(sh)
    if f:
        d.rectangle([x, y, x + w, y + h], fill=f)
    if sh.has_table:
        t = sh.table
        cy = y
        for r in t.rows:
            cx = x
            for ci, c in enumerate(r.cells):
                cw, chh = px(t.columns[ci].width), px(r.height)
                cf = None
                try:
                    if c.fill.type == 1:
                        cf = "#%06X" % int(str(c.fill.fore_color.rgb), 16)
                except Exception:
                    pass
                d.rectangle([cx, cy, cx + cw, cy + chh], fill=cf,
                            outline="#DDD")
                runs = [rr for p in c.text_frame.paragraphs for rr in p.runs]
                if runs:
                    sz = runs[0].font.size.pt if runs[0].font.size else 11
                    col = "#000"
                    try:
                        col = "#%06X" % int(str(runs[0].font.color.rgb), 16)
                    except Exception:
                        pass
                    d.text((cx + 4, cy + 3), c.text[:60], font=font(sz),
                           fill=col)
                cx += cw
            cy += px(r.height)
        return
    if not sh.has_text_frame or not sh.text_frame.text.strip():
        return
    ty = y + 2
    for p in sh.text_frame.paragraphs:
        txt = "".join(rr.text for rr in p.runs)
        if not txt:
            ty += 8
            continue
        sz = max([rr.font.size.pt for rr in p.runs if rr.font.size] or [12.7])
        col = "#232323"
        for rr in p.runs:
            try:
                col = "#%06X" % int(str(rr.font.color.rgb), 16)
                break
            except Exception:
                pass
        fo = font(sz)
        avail = w - 4
        line, cur = "", 0
        for ch in txt:
            cw = d.textlength(ch, font=fo)
            if cur + cw > avail and line:
                d.text((x + 3, ty), line, font=fo, fill=col)
                ty += int(sz * 1.2 * SCALE / 72)
                line, cur = "", 0
            line += ch
            cur += cw
        if line:
            d.text((x + 3, ty), line, font=fo, fill=col)
            ty += int(sz * 1.2 * SCALE / 72)


def main(path, pages, out_prefix):
    prs = Presentation(path)
    W, H = px(prs.slide_width), px(prs.slide_height)
    for n in pages:
        s = prs.slides[n - 1]
        im = Image.new("RGB", (W, H), "white")
        d = ImageDraw.Draw(im)
        for sh in s.shapes:
            draw_shape(d, im, sh)
        d.rectangle([0, 0, W - 1, H - 1], outline="#BBB")
        im.save("%s_p%02d.png" % (out_prefix, n))
        print("%s_p%02d.png" % (out_prefix, n))


if __name__ == "__main__":
    main(sys.argv[1], [int(x) for x in sys.argv[2].split(",")], sys.argv[3])
