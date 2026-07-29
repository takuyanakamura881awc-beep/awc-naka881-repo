# -*- coding: utf-8 -*-
"""生成したPPTXの版面チェック（溢れ・はみ出し・コントラスト）。

この環境ではLibreOfficeでのレンダリングができないため、
座標と文字幅の計算で検証する。
"""
import sys
import unicodedata

from pptx import Presentation
from pptx.util import Emu, Pt

EMU_IN = 914400


def chw(ch, size):
    return size if unicodedata.east_asian_width(ch) in "WFA" else size * 0.55


def wrap_lines(text, avail_pt, size):
    """折り返し後の行数を概算"""
    if not text:
        return 0
    n, cur = 1, 0.0
    for ch in text:
        w = chw(ch, size)
        if cur + w > avail_pt and cur > 0:
            n += 1
            cur = 0.0
        cur += w
    return n


def lum(rgb):
    c = [(rgb >> 16 & 255) / 255, (rgb >> 8 & 255) / 255, (rgb & 255) / 255]
    c = [x / 12.92 if x <= 0.03928 else ((x + 0.055) / 1.055) ** 2.4 for x in c]
    return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]


def contrast(a, b):
    la, lb = lum(a), lum(b)
    return (max(la, lb) + 0.05) / (min(la, lb) + 0.05)


def shape_fill(sh):
    try:
        if sh.fill.type is not None and sh.fill.type == 1:
            return int(str(sh.fill.fore_color.rgb), 16)
    except Exception:
        pass
    return None


def rect(sh):
    return (sh.left, sh.top, sh.left + sh.width, sh.top + sh.height)


def main(path):
    prs = Presentation(path)
    SW, SH = prs.slide_width, prs.slide_height
    overflow, outside, lowc = [], [], []

    for si, s in enumerate(prs.slides, 1):
        boxes = [(rect(sh), shape_fill(sh)) for sh in s.shapes
                 if shape_fill(sh) is not None]
        for sh in s.shapes:
            if sh.left is None:
                continue
            l, t, r, b = rect(sh)
            if l < -1000 or t < -1000 or r > SW + 1000 or b > SH + 1000:
                outside.append((si, sh.name,
                                round(Emu(l).inches, 2), round(Emu(t).inches, 2),
                                round(Emu(r).inches, 2), round(Emu(b).inches, 2)))
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            tf = sh.text_frame
            ml = tf.margin_left or 0
            mr = tf.margin_right or 0
            mt = tf.margin_top or 0
            mb = tf.margin_bottom or 0
            avail_pt = Emu(sh.width - ml - mr).inches * 72
            need = 0.0
            maxsz = 0
            for para in tf.paragraphs:
                txt = "".join(rr.text for rr in para.runs)
                sz = max([rr.font.size.pt for rr in para.runs
                          if rr.font.size] or [12.7])
                maxsz = max(maxsz, sz)
                ls = para.line_spacing or 1.0
                if not isinstance(ls, float) and not isinstance(ls, int):
                    ls = 1.0
                n = max(1, wrap_lines(txt, avail_pt, sz))
                sb = para.space_before.pt if para.space_before else 0
                sa = para.space_after.pt if para.space_after else 0
                need += n * sz * 1.2 * ls + sb + sa
            box_pt = Emu(sh.height - mt - mb).inches * 72
            if need > box_pt + 1.0:
                overflow.append((si, sh.name,
                                 sh.text_frame.text.replace("\n", "|")[:34],
                                 round(need, 1), round(box_pt, 1)))
            # コントラスト（自身の塗り、なければ重なる矩形の塗り、なければ白）
            bg = shape_fill(sh)
            if bg is None:
                cx, cy = (l + r) // 2, (t + b) // 2
                cands = [f for (bl, bt, br, bb), f in boxes
                         if bl <= cx <= br and bt <= cy <= bb]
                bg = cands[-1] if cands else 0xFFFFFF
            for para in tf.paragraphs:
                for rr in para.runs:
                    if not rr.text.strip():
                        continue
                    try:
                        fg = int(str(rr.font.color.rgb), 16)
                    except Exception:
                        continue
                    sz = rr.font.size.pt if rr.font.size else 12.7
                    need_ratio = 3.0 if (sz >= 18 or (sz >= 14 and rr.font.bold)) \
                        else 4.5
                    cr = contrast(fg, bg)
                    if cr < need_ratio:
                        lowc.append((si, sh.name, rr.text[:22],
                                     "#%06X" % fg, "#%06X" % bg,
                                     round(cr, 2), sz))

    print("== 溢れ（推定） ==", len(overflow))
    for o in overflow:
        print("  p%-3d %-22s %-36s need=%.1fpt box=%.1fpt" % o)
    print("== スライド外 ==", len(outside))
    for o in outside:
        print("  ", o)
    print("== コントラスト未達 ==", len(lowc))
    for o in lowc:
        print("  p%-3d %-20s %-24s fg=%s bg=%s %.2f:1 %.1fpt" % o)
    return len(overflow) + len(outside) + len(lowc)


if __name__ == "__main__":
    sys.exit(0 if main(sys.argv[1]) == 0 else 1)
